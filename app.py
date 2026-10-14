"""
Store Signage Schedule Extractor
--------------------------------
Run:   streamlit run app.py
Needs: pip install streamlit pdfplumber pypdfium2 pandas openpyxl pillow reportlab anthropic
"""
import base64
import hashlib
import io
import json
import math
import os
import re
import zipfile
from collections import Counter, defaultdict

import difflib

import pandas as pd
import pdfplumber
import pypdfium2 as pdfium
import streamlit as st
from PIL import Image, ImageDraw

# --------------------------------------------------------------------------------------
# STREAMLIT CONFIG (Must be the very first st.* command!)
# --------------------------------------------------------------------------------------
st.set_page_config(page_title="Store Signage Schedule Extractor", layout="wide")

# --------------------------------------------------------------------------------------
# Patterns
# --------------------------------------------------------------------------------------
CODE_PATTERN = r"(?:TK\d{2}(?:-\d{1,3}[A-Z]{0,2})?(?:-[A-Z]{1,2})?|TKM-\d+[A-Z]*|TK_[A-Za-z]+)"
ROW_RE = re.compile(rf"^(L\d{{2}})\s+(.+?)\s+({CODE_PATTERN})\s+(\d+)(?=\s|MOD|$)(.*)$")   # count can run straight into 'MOD6' ('1MOD6')


# --------------------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------------------
TK_LEVEL_MAP = {"GF": "L00", "G": "L00", "FF": "L01", "SF": "L02", "TF": "L03", "LG": "B01", "LGF": "B01", "BF": "B01", "B": "B01", "MEZZ": "M00"}


def tk_table_rows_words(words, vrules, page_no):
    """Read the TK_Signage table by its columns (Level / Family and Type / TK Code / Count / MOD) instead of by text line.
    Copes with older schedules: 'GF'/'FF' levels, revision notes printed alongside the table, codes like TK007-A or TK402-01_DS."""
    lines = defaultdict(list)
    for w in words:
        lines[round(w["top"] / 2.5)].append(w)
    hdr = None
    for k in sorted(lines):
        ws = sorted(lines[k], key=lambda w: w["x0"])
        txt = [w["text"] for w in ws]
        if "Level" in txt and "Count" in txt and any(t == "Family" for t in txt) and "Code" in txt:
            lv = ws[txt.index("Level")]
            fam = ws[txt.index("Family")]
            ci = txt.index("Code")
            code = ws[ci - 1] if ci and txt[ci - 1] == "TK" else ws[ci]
            cnt = ws[txt.index("Count")]
            mod = next((w for w in ws if w["text"] == "MOD" and w["x0"] > cnt["x0"]), None)
            top, bot = min(w["top"] for w in ws), max(w["bottom"] for w in ws)
            heads = [(lv["x0"] + lv["x1"]) / 2, (fam["x0"] + fam["x1"]) / 2, (code["x0"] + ws[ci]["x1"]) / 2, (cnt["x0"] + cnt["x1"]) / 2]
            if mod:
                heads.append((mod["x0"] + mod["x1"]) / 2)
            # column rules: vertical edges crossing the header row
            vx = sorted({round(x, 1) for x, t0, t1 in vrules if t0 <= top + 2 and t1 >= bot - 2 and heads[0] - 200 < x < heads[-1] + 200})
            def bounds(c, i):
                left = max((x for x in vx if x < c), default=None)
                right = min((x for x in vx if x > c), default=None)
                if left is None:
                    left = (heads[i - 1] + c) / 2 if i else c - 60
                if right is None:
                    right = (heads[i + 1] + c) / 2 if i + 1 < len(heads) else c + 40
                return left, right
            b_lv, b_fam, b_code, b_cnt = (bounds(heads[k], k) for k in range(4))
            b_mod = bounds(heads[4], 4) if mod else None
            hdr = {"top": bot, "lv": b_lv[0] - 2, "fam": b_fam[0] - 1, "code": b_code[0] - 1, "cnt": b_cnt[0] - 1,
                   "mod": (b_mod[0] - 1) if b_mod else None, "right": (b_mod[1] if b_mod else b_cnt[1]) + 1}
            break
    if not hdr:
        return []
    rows, cur = [], None
    for k in sorted(lines):
        ws = sorted((w for w in lines[k] if w["top"] > hdr["top"] and hdr["lv"] <= w["x0"] < hdr["right"]), key=lambda w: w["x0"])
        if not ws:
            continue
        lvw = [w["text"] for w in ws if w["x0"] < hdr["fam"]]
        famw = [w["text"] for w in ws if hdr["fam"] <= w["x0"] < hdr["code"]]
        rest = [w for w in ws if w["x0"] >= hdr["code"] and (hdr["mod"] is None or w["x0"] < hdr["mod"])]
        modw = [w["text"] for w in ws if hdr["mod"] is not None and w["x0"] >= hdr["mod"]]
        nums = [w for w in rest if re.fullmatch(r"\d{1,4}", w["text"])]
        qty = nums[-1] if nums else None
        codew = [w["text"] for w in rest if w is not qty and w["x0"] < hdr["cnt"]]
        level = " ".join(lvw).upper().replace("_", " ").strip() if lvw else ""
        l0 = level.split(" ")[0] if level else ""
        if l0 and re.fullmatch(r"L\d{2}|B\d{2}|M\d{2}|[A-Z]{1,4}", l0) and codew and qty is not None:
            lev = l0 if re.fullmatch(r"[LBM]\d{2}", l0) else TK_LEVEL_MAP.get(l0, l0)
            if " " in level and not re.search(r"plan|fixture", level, re.I):
                lev += " " + level.split(" ", 1)[1].title()          # 'GF_Upper Floor' -> 'L00 Upper Floor'; 'L00_Fixture Plan' -> 'L00'
            cur = {"level": lev, "code": " ".join(codew).strip(), "desc": " ".join(famw), "qty": int(qty["text"]),
                   "mod": next((m for m in modw if re.fullmatch(r"MOD\d+", m)), ""), "sched_page": page_no}
            rows.append(cur)
        elif cur is not None and famw and not lvw and not codew and qty is None:
            cur["desc"] += " " + " ".join(famw)             # description wrapped onto a second line
        elif not lvw and qty is not None and not famw and not codew:
            cur = None                                       # the total under the Count column - end of the table
    for r in rows:
        r["desc"] = clean_description(r["desc"])
    return [r for r in rows if r["code"] and r["code"].upper().startswith("TK")
            and not re.search(r"digital\s*screen", r["code"] + " " + r["desc"], re.I)]      # screens are not signage


def tk_table_rows(plumber_page, page_no):
    words = plumber_page.extract_words(keep_blank_chars=False, use_text_flow=False, x_tolerance=1.5, y_tolerance=2)
    vr = [(e["x0"], e["top"], e["bottom"]) for e in plumber_page.edges if abs(e["x0"] - e["x1"]) < 1]
    return tk_table_rows_words(words, vr, page_no)


def tk_table_rows_pdfium(page, page_no):
    """Same table reader for rotated sheets: words rebuilt from characters in displayed coordinates, rules from the drawing's paths."""
    import pypdfium2.raw as pdfium_c
    tf, dw, dh = geom(page)
    tp = page.get_textpage()
    chars = []
    try:
        for i in range(tp.count_chars()):
            c = tp.get_text_range(i, 1)
            if not c or c in "\r\n":
                chars.append(None)
                continue
            x0, y0, x1, y1 = box_tf(tf, *tp.get_charbox(i))
            chars.append((c, x0, y0, x1, y1, pdfium_c.FPDFText_GetFontSize(tp.raw, i)))
    finally:
        tp.close()
    words, cur = [], None
    for ch in chars + [None]:
        if ch is None or ch[0] == " ":
            if cur:
                words.append(cur); cur = None
            continue
        c, x0, y0, x1, y1, fs = ch
        if c == "_" and cur:                                  # underscore glyph boxes sit low / off-line: keep them in the word
            cur["text"] += c
            cur["us"] = True
            continue
        fs = max(fs, 1.0)
        lim = 1.6 * fs if (cur and cur.get("us")) else 0.45 * fs
        if cur and abs(cur["base"] - y1) < 0.45 * fs and -1.0 <= x0 - cur["x1"] < lim:
            cur["us"] = False
            cur["text"] += c; cur["x1"] = max(cur["x1"], x1); cur["top"] = min(cur["top"], y0); cur["bottom"] = max(cur["bottom"], y1)
        else:
            if cur:
                words.append(cur)
            cur = {"text": c, "x0": x0, "x1": x1, "top": y0, "bottom": y1, "base": y1}
    vr = []
    for o in page.get_objects(max_depth=15):
        if o.type == pdfium_c.FPDF_PAGEOBJ_PATH:
            x0, y0, x1, y1 = box_tf(tf, *o.get_bounds())
            if x1 - x0 < 1.5 and y1 - y0 > 5:
                vr.append(((x0 + x1) / 2, y0, y1))
    return tk_table_rows_words(words, vr, page_no)


def is_schedule_page(text: str) -> bool:
    t = text.lower()
    return "tk code" in t and "family and type" in t


def classify_page(text: str, sheet_no: str = ""):
    t = text.lower()
    m = re.search(r"\b(L\d{2})[_ ]Signage Plan", text, re.I)
    if m:
        return "plan", m.group(1).upper()
    if any(kw in t for kw in ["elevation", "elev", "internal elevation", "backwall", "shopfront", "section"]) or re.match(r"^A[45]\d{2}", sheet_no):
        return "elevation", ""
    return "other", ""


def clean_description(raw: str) -> str:
    family, _, typ = raw.partition(": ")
    family, typ = family.strip(), typ.strip()
    if not typ or typ == family:
        return family
    return f"{family} - {typ}"


def size_from_name(desc: str) -> str:
    out = []
    m = re.search(r"(\d{3,4})\s*[x×]\s*(\d{2,4})", desc)
    if m:
        out.append(f"{m.group(1)} x {m.group(2)}")
    m = re.search(r"\b(\d{2,3})mm\b", desc)
    if m:
        out.append(f"{m.group(1)}mm")
    m = re.search(r"(?:^|_)(\d{2,3})_", desc)
    if m and not out:
        out.append(f"{m.group(1)}mm letter height (name)")
    if "roundel" in desc.lower():
        m = re.search(r"(\d{3,4})\s*$", desc)
        if m:
            out.append(f"{m.group(1)} roundel (name)")
    return "; ".join(out)


# --------------------------------------------------------------------------------------
# PDF layer (pypdfium2) - rotation-safe text positions and crops
# Many drawing sheets are stored with a /Rotate flag. pdfplumber reads those as mirrored /
# vertical text and its crops land in the wrong place, so positions come from here instead.
# All coordinates below are "displayed" points: origin top-left, y down.
# --------------------------------------------------------------------------------------
def geom(page):
    rot = page.get_rotation() % 360
    L, B, R, T = page.get_bbox()
    w, h = R - L, T - B
    dw, dh = (w, h) if rot in (0, 180) else (h, w)

    def tf(x, y):
        u, v = x - L, y - B
        if rot == 0:
            return u, h - v
        if rot == 90:
            return v, u
        if rot == 180:
            return w - u, v
        return h - v, w - u

    return tf, dw, dh


def box_tf(tf, l, b, r, t):
    xs, ys = zip(tf(l, b), tf(r, t), tf(l, t), tf(r, b))
    return min(xs), min(ys), max(xs), max(ys)


def page_runs(page):
    """Text runs (roughly one per line/label) with displayed bboxes."""
    tf, dw, dh = geom(page)
    tp = page.get_textpage()
    out = []
    try:
        for i in range(tp.count_rects()):
            l, b, r, t = tp.get_rect(i)
            s = tp.get_text_bounded(l, b, r, t)
            s = re.sub(r"\s+", " ", s.replace("\ufffe", "-").replace("\x02", "-")).strip()
            if not s:
                continue
            x0, y0, x1, y1 = box_tf(tf, l, b, r, t)
            out.append({"t": s, "x0": x0, "y0": y0, "x1": x1, "y1": y1, "h": max(y1 - y0, 1.0)})
    finally:
        tp.close()
    return out, dw, dh


def runs_to_text(runs):
    """Rebuild reading-order lines from runs (used for rotated pages, where pdfplumber text is unusable)."""
    if not runs:
        return ""
    runs = sorted(runs, key=lambda r: ((r["y0"] + r["y1"]) / 2, r["x0"]))
    lines, cur, cy = [], [], None
    for r in runs:
        c = (r["y0"] + r["y1"]) / 2
        if cur and abs(c - cy) <= 0.6 * r["h"]:
            cur.append(r)
        else:
            if cur:
                lines.append(cur)
            cur, cy = [r], c
    if cur:
        lines.append(cur)
    return "\n".join(" ".join(r["t"] for r in sorted(l, key=lambda r: r["x0"])) for l in lines)


def crop_image(page, box, scale, marks=()):
    tf, dw, dh = geom(page)
    x0, y0, x1, y1 = box
    im = page.render(scale=scale, crop=(x0, dh - y1, dw - x1, y0)).to_pil().convert("RGB")
    d = ImageDraw.Draw(im)
    for m in marks:
        d.rectangle([(m[0] - x0) * scale - 4, (m[1] - y0) * scale - 4,
                     (m[2] - x0) * scale + 4, (m[3] - y0) * scale + 4], outline=(220, 0, 0), width=3)
    return im


def crop_png(page, centre, half_w, half_h, scale, marks=()):
    tf, dw, dh = geom(page)
    cx, cy = centre
    box = (max(0, cx - half_w), max(0, cy - half_h), min(dw, cx + half_w), min(dh, cy + half_h))
    im = crop_image(page, box, scale, marks)
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


_DASHES = str.maketrans({"\u2010": "-", "\u2011": "-", "\u2012": "-", "\u2013": "-", "\u2014": "-",
                         "\u2212": "-", "\u00ad": "-", "\ufffe": "-", "\u00a0": " "})


def find_code_boxes(page, code_re, code_map):
    """[(code, (x0,y0,x1,y1) displayed)] for every exact occurrence of a schedule code on the page."""
    if code_re is None:
        return []
    tf, dw, dh = geom(page)
    tp = page.get_textpage()
    out = []
    try:
        raw = tp.get_text_range()
        norm = raw.translate(_DASHES).upper()
        ok = len(raw) == tp.count_chars()
        runs = None
        for m in code_re.finditer(norm):
            code = code_map[m.group(1)]
            box = None
            if ok:
                try:
                    bs = [tp.get_charbox(k) for k in range(m.start(), m.end()) if not raw[k].isspace()]
                    if bs:
                        box = box_tf(tf, min(b[0] for b in bs), min(b[1] for b in bs),
                                     max(b[2] for b in bs), max(b[3] for b in bs))
                except Exception:
                    box = None
            if box is None:                      # fall back to the whole text run containing the code
                if runs is None:
                    runs = page_runs(page)[0]
                r = next((r for r in runs if m.group(1) in r["t"].translate(_DASHES).upper()), None)
                if r:
                    box = (r["x0"], r["y0"], r["x1"], r["y1"])
            if box:
                out.append((code, box))
    finally:
        tp.close()
    return out


def nearest_view_title(runs, hit) -> str:
    cands = [r for r in runs if "elevation" in r["t"].lower() or re.match(r"^L\d{2}_", r["t"])]
    if not cands:
        return ""
    cx, cy = (hit[0] + hit[2]) / 2, (hit[1] + hit[3]) / 2

    def dist(r):
        dy = (r["y0"] + r["y1"]) / 2 - cy
        return math.hypot((r["x0"] + r["x1"]) / 2 - cx, dy * 3 if dy < 0 else dy)

    return min(cands, key=dist)["t"][:70]


# --------------------------------------------------------------------------------------
# TK Maxx analysis
# --------------------------------------------------------------------------------------
def _analyse(pdf_bytes, elev_n, incl_plan, ehw, ehh, phw, phh, dpi):
    _reader_version = "2026-10-14-fast"   # changes the cache key of this function so results stored before the schedule-row fix are not reused
    rows, unmatched, notes, page_rows, warnings = [], [], [], [], []
    doc = pdfium.PdfDocument(pdf_bytes)
    n_pages = len(doc)
    scale = dpi / 72.0

    # ---- text per page. Upright pages: pdfplumber layout text (as before). Rotated pages: rebuilt from pdfium.
    # pdfplumber's layout text is slow on dense drawings (90s+ a sheet on some packs), so it is only used for the schedule page(s);
    # every other page is read with pdfium, which is what classification and code-finding need.
    texts, rotated = {}, set()
    upright_sched = []
    for i in range(1, n_pages + 1):
        pg = doc[i - 1]
        try:
            if pg.get_rotation() % 360:
                rotated.add(i)
                texts[i] = runs_to_text(page_runs(pg)[0])
            else:
                tp = pg.get_textpage()
                try:
                    texts[i] = tp.get_text_range().replace("\r\n", "\n").replace("\r", "\n")
                finally:
                    tp.close()
                if is_schedule_page(texts[i]):
                    upright_sched.append(i)
        except Exception:
            texts[i] = ""
        finally:
            pg.close()
    if upright_sched:
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            for i in upright_sched:
                try:
                    texts[i] = pdf.pages[i - 1].extract_text() or texts[i]
                except Exception:
                    pass

    # ---- schedule rows
    row_pages = Counter()
    for i, t in texts.items():
        if is_schedule_page(t):
            for line in t.splitlines():
                if ROW_RE.match(line.strip()):
                    row_pages[i] += 1
    sched_pages_all = [i for i in texts if is_schedule_page(texts[i])]
    sched_pages = [i for i in sched_pages_all if row_pages[i] >= 1]
    table_rows = {}
    fallback = not sched_pages
    scan = list(texts) if fallback else sched_pages
    for i in sched_pages_all:
        try:
            if i in rotated:
                _pg = doc[i - 1]
                try:
                    trows = tk_table_rows_pdfium(_pg, i)
                finally:
                    _pg.close()
            else:
                with pdfplumber.open(io.BytesIO(pdf_bytes)) as _pl:
                    trows = tk_table_rows(_pl.pages[i - 1], i)
        except Exception:
            trows = []
        if len(trows) > row_pages.get(i, 0):
            table_rows[i] = trows
    if table_rows and fallback:
        fallback = False
        sched_pages = sorted(table_rows)
        scan = sched_pages
    for i in sorted(table_rows):
        rows.extend(table_rows[i])
        if i not in sched_pages:
            sched_pages = sorted(set(sched_pages) | {i})
    for i in scan:
        if i in table_rows:
            continue
        for line in texts[i].splitlines():
            line = line.strip()
            m = ROW_RE.match(line)
            if m:
                level, raw, code, qty, tail = m.groups()
                mod = re.search(r"\bMOD\d+\b", tail)
                rows.append({"level": level, "code": code, "desc": clean_description(raw),
                             "qty": int(qty), "mod": mod.group(0) if mod else "", "sched_page": i})
            elif not fallback and re.match(r"^L\d{2}\s", line) and "TK" in line:
                unmatched.append(f"p{i}: {line}")

    for i, t in texts.items():
        for line in t.splitlines():
            if re.search(r"\bFFL\b", line) and line.strip() not in notes:
                notes.append(line.strip())

    codes = {r["code"] for r in rows}
    code_map = {c.upper(): c for c in codes}
    code_re = (re.compile(r"(?<![A-Z0-9_-])(" + "|".join(re.escape(c) for c in sorted(code_map, key=len, reverse=True)) +
                          r")(?![A-Z0-9_-])") if code_map else None)

    # ---- find every code on every drawing page
    hits = defaultdict(list)
    runs_cache = {}
    for i in range(1, n_pages + 1):
        c = Counter(re.findall(r"\bA\d{3}\b", texts[i]))
        sheet = c.most_common(1)[0][0] if c else ""
        try:                                   # the sheet's own number is the big 'A415' in the title block, not the most-quoted one
            _pg = doc[i - 1]
            _runs, _dw, _dh = page_runs(_pg)
            _pg.close()
            _tb = [r for r in _runs if r["x0"] > 0.8 * _dw and r["y0"] > 0.8 * _dh and re.match(r"^A\d{3}\b", r["t"].strip())]
            if _tb:
                sheet = re.match(r"^A\d{3}", max(_tb, key=lambda r: r["h"])["t"].strip()).group(0)
        except Exception:
            pass
        is_sched = i in sched_pages
        kind, plevel = classify_page(texts[i], sheet) if not is_sched else ("schedule", "")
        n_here = 0
        if not is_sched and code_re is not None:
            pg = doc[i - 1]
            try:
                found = find_code_boxes(pg, code_re, code_map)
                if found and kind == "other":
                    kind = "elevation"
                if found:
                    runs = page_runs(pg)[0] if kind == "elevation" else []
                for code, box in found:
                    n_here += 1
                    hits[code].append({"page": i, "kind": kind, "level": plevel, "sheet": sheet, "box": box,
                                       "view": nearest_view_title(runs, box) if kind == "elevation" else ""})
            finally:
                pg.close()
        page_rows.append({"page": i, "type": kind, "level": plevel, "sheet": sheet, "rotated": i in rotated,
                          "chars": len(texts[i]), "codes_found": n_here})

    blank = [p["page"] for p in page_rows if p["chars"] < 20 and p["type"] != "schedule"]
    if blank:
        warnings.append("No readable text on page(s) " + ", ".join(f"p{p}" for p in blank[:20]) +
                        " (scanned or outlined drawings). Signs on those pages can't be found automatically.")

    # ---- rows, snippets, checks
    snips, out_rows = {}, []
    page_dims = {}
    for r in rows:
        code, level = r["code"], r["level"]
        elev = [h for h in hits[code] if h["kind"] == "elevation"]
        plan = [h for h in hits[code] if h["kind"] == "plan" and h["level"] == level]
        other_plan = [h for h in hits[code] if h["kind"] == "plan" and h["level"] != level]
        pages_found = sorted({(h["page"], h["kind"], h["sheet"] or h["level"]) for h in hits[code]})

        chosen = [(h, "elevation") for h in elev[:elev_n]]
        if incl_plan and plan:
            chosen.append((plan[0], "plan"))
        if not chosen and hits[code]:             # code exists on a drawing we couldn't classify - still show it
            chosen = [(hits[code][0], hits[code][0]["kind"])]
        lst = []
        for h, k in chosen:
            pg = doc[h["page"] - 1]
            try:
                hw, hh = (ehw, ehh) if k == "elevation" else (phw, phh)
                cx, cy = (h["box"][0] + h["box"][2]) / 2, (h["box"][1] + h["box"][3]) / 2
                png = crop_png(pg, (cx, cy), hw, hh, scale, [h["box"]])
            finally:
                pg.close()
            lst.append({"kind": k, "page": h["page"], "sheet": h["sheet"], "view": h["view"],
                        "label": f"{k} p{h['page']} {h['sheet']} {h['view']}".strip(), "png": png})
        snips[f"{level}|{code}"] = lst

        pos = []
        if plan:
            pos.append(f"{level} plan")
        seen = set()
        for h in elev:
            label = f"{h['sheet']} {h['view']}".strip()
            if label and label not in seen:
                seen.add(label)
                pos.append(label)
        n_plan = len(plan)
        if n_plan == 0:
            plan_check = "not tagged on plan"
        elif n_plan == r["qty"]:
            plan_check = "OK"
        else:
            plan_check = f"plan shows {n_plan}"
        issues = []
        if not hits[code]:
            issues.append("code not found on any drawing")
        elif not elev:
            issues.append("no elevation found" + (" (found on plan only)" if plan or other_plan else ""))
        if n_plan and n_plan != r["qty"]:
            issues.append("qty differs from plan")
        out_rows.append({**r, "size_name": size_from_name(r["desc"]), "position": "; ".join(pos[:4]),
                         "plan_check": plan_check, "check": "; ".join(issues),
                         "pages": ", ".join(f"p{p} {k} {s}".strip() for p, k, s in pages_found)})
    doc.close()
    return {
        "rows": out_rows, "snips": snips, "unmatched": unmatched, "notes": notes, "warnings": warnings,
        "pages": page_rows, "schedule_pages": sched_pages, "fallback": fallback,
        "hits": [{"code": c, **{k: v for k, v in h.items() if k in ("page", "kind", "level", "sheet", "view")}}
                 for c, hl in hits.items() for h in hl],
    }


analyse = st.cache_data(show_spinner="Reading drawing pack...")(_analyse)


# --------------------------------------------------------------------------------------
# M&S analysis (client drawing binder: signs are red callouts on plans and elevations)
# --------------------------------------------------------------------------------------
MS_SIGN_RE = re.compile(
    r"\bSIGN(?:S|AGE)?\b|LIGHT ?BOX|\bVINYLS?\b|\bUPC\b|ARTWORK REF|\bFASCIA\b|\bGOBO\b|MANIFESTATION|DIRECTORY|TOTEM|"
    r"QUALIFI?ER|LETTERING|NOTICE ?BOARD|HISTORY WALL|\bBANNER\b|ROUNDEL|NAVIGATIONAL|PUSH THROUGH|\bHALO\b|"
    r"\bHUB\b|WELCOME WALL|MEDIA SCREEN", re.I)
MS_NOISE_RE = re.compile(
    r"COPYRIGHT|DO NOT SCALE|REPRODUCED|ALL DIMENSIONS|FIRE (?:DOOR|ACTION|EXTINGUISHER|ESCAPE|STRATEGY)|ISO ?\d{4}|"
    r"EMERGENCY|EXIT SIGN|SIGNS AND SIGNALS|WARNING SIGN|SIGN CODE|KEY COMMENTS|UPDATED IN ACCORDANCE|"
    r"REVISED CONSTRUCTION|ISSUED AS PART|SIGN HEIGHT|MANDATORY|LUMINOUS|ASBESTOS|PLEASE REFER TO ELEVATION|"
    r"SIGNAGE KEY|DESCRIPTION CHECKED|TENDER ISSUE|CONSTRUCTION ISSUE|FIRST ISSUE|SIGN\s*[:=]|RELOCATED TO AVOID|"
    r"MEDIA SCREEN|^(?:GROUND|FIRST|BASEMENT) FLOOR\b|SIGNED OFF|FIRE HOSE|QUALIFIER DETAIL|DETAIL\.? ?\d? ?: ?\d+|"
    r"MERCHANDISER|TILL ASSISTED|ASSISTED (?:WRAP|TILL)|ACC UNIT|TILL END|CARD/CARD|CARD AND CASH TILL|\bUNIT\b.{0,25}UPC|"
    r"NOT TO EXCEED|COVED|SKIRTING|PLASTERBOARD|SOFFIT|FLOOR TRANSITION|MCCUE|BUMP RAIL|POWER FOR PHONE|POWER TO BE HIDDEN|SIZE TO BE CONFIRMED ON SITE\)? POWER", re.I)
MS_LEGEND_RE = re.compile(
    r"^(?:ILLUMINATED SIGN|NON-ILLUMINATED SIGN / ARTWORK|LIGHTBOX|MID FLOOR (?:ILLUMINATED|DIRECTIONAL) SIGNS?.*|"
    r"SUPPLEMENTARY SIGNAGE|DOOR VINYLS ARTWORK DETAILS.*|EXISTING MANIFESTATIONS REPLACED.*|EXAMPLE - MID FLOOR SIGNAGE)$", re.I)
MS_STOP = set("A AN AND AS AT BE BY FOR FROM IN IS IT OF ON OR TO THE WITH ALL NEW REFER DRAWING DETAILS FURTHER SIGN SIGNAGE "
              "ILLUMINATED SIGNS MOUNTED WALL ARTWORK REF AI TBC PLAN UPC APPLIED GRAPHICS".split())
LEVEL_NAMES = {"B": "Basement", "GF": "Ground floor", "FF": "First floor"}


def sentence_case(t):
    if t.isupper():
        t = t.lower().capitalize()
        t = re.sub(r"\b(upc|affl|led|leds|bu|ffl|ai|eas|ral)\b", lambda m: m.group(1).upper(), t)
        t = t.replace("m&s", "M&S")
    return t


def ms_cluster(runs):
    runs = [r for r in runs if not (r["y1"] - r["y0"] > 1.6 * (r["x1"] - r["x0"]) and len(r["t"]) > 3)]  # skip vertical text
    n = len(runs)
    par = list(range(n))

    def f(a):
        while par[a] != a:
            par[a] = par[par[a]]
            a = par[a]
        return a

    order = sorted(range(n), key=lambda i: runs[i]["y0"])
    for ii, i in enumerate(order):
        a = runs[i]
        for j in order[ii + 1: ii + 70]:
            b = runs[j]
            h = max(a["h"], b["h"])
            if b["y0"] - a["y1"] > 0.9 * h:
                break
            if abs(a["h"] - b["h"]) > 0.45 * h:
                continue
            gap = b["y0"] - a["y1"]
            if gap < -0.6 * h:
                continue
            if (abs(a["x0"] - b["x0"]) < 1.2 * h or abs((a["x0"] + a["x1"]) / 2 - (b["x0"] + b["x1"]) / 2) < 1.2 * h
                    or abs(a["x1"] - b["x1"]) < 1.2 * h):
                par[f(i)] = f(j)
    groups = defaultdict(list)
    for i in range(n):
        groups[f(i)].append(runs[i])
    out = []
    for g in groups.values():
        g.sort(key=lambda r: (r["y0"], r["x0"]))
        out.append({"text": re.sub(r"\s+", " ", runs_to_text(g).replace("\n", " ")).strip(),
                    "x0": min(r["x0"] for r in g), "y0": min(r["y0"] for r in g),
                    "x1": max(r["x1"] for r in g), "y1": max(r["y1"] for r in g), "h": g[0]["h"]})
    return out


def ms_title(runs, dw, dh):
    """Drawing title + number from the title block (bottom-right of the sheet)."""
    tb = [r for r in runs if r["x0"] > 0.55 * dw and r["y0"] > 0.62 * dh]
    cl = ms_cluster(tb)
    bad = re.compile(r"COPYRIGHT|REPRODUCED|HALLIDAY|CONSTRUCTION|DUNDRUM|IRELAND|NOTES|ARCHITECTURE|ISSUE|KEY COMMENTS|"
                     r"UPDATED|DESCRIPTION|PROJECT|CONSULTANT|REVISED|AS INDICATED|^[A-Z]{3,4} ?\d{2}$|\d{2}\.\d{2}\.\d{2}", re.I)
    kw = re.compile(r"signage|elevation|plan|details|layout|finishes|schedule|setting|section|till|door", re.I)
    c = [x for x in cl if 6 < len(x["text"]) <= 72 and not bad.search(x["text"]) and kw.search(x["text"])]
    prim = re.compile(r"signage|elevation|plan\b|details|layout|setting|section", re.I)    # a sheet title beats a legend heading ('Wall Finishes')
    title = max(c, key=lambda x: (bool(prim.search(x["text"])), x["h"]))["text"] if c else ""
    title = re.sub(r"^(?:DRAWING\s+TITLE|TITLE)\s*:?\s*", "", title, flags=re.I)
    num = ""
    dums = [r for r in tb if re.fullmatch(r"DUM\d+ ?[A-Z]?", r["t"])]
    nums = [r for r in tb if re.fullmatch(r"\d{5}", r["t"])]
    if nums:
        ref = dums[0] if dums else nums[0]
        num = [r["t"] for r in sorted(nums, key=lambda r: math.hypot(r["x0"] - ref["x0"], r["y0"] - ref["y0"]))]
    return title, num
    return title, num


def ms_level(title, text):
    t = (title or "").lower()
    for key, lv in (("basement", "B"), ("boh", "B"), ("ground", "GF"), ("foodhall", "GF"), ("food hall", "GF"), ("first", "FF")):
        if key in t:
            return lv
    if re.match(r"^ff\b", t):
        return "FF"
    if re.match(r"^gf\b", t):
        return "GF"
    return ""


def ms_page_kind(title, text):
    t = (title or "").lower()
    if "signage" in t or "signage key" in text.lower():
        return "plan"
    if re.search(r"elevation|till details|till hanger|section|setting out|manifestation|\bdetails\b", t):
        return "elevation"
    return "other"


def ms_parse(text):
    t = text
    refs = []
    for m in re.finditer(r"UPC:?\s*(\d{6,8})", t, re.I):
        refs.append(f"UPC {m.group(1)}")
    art = None
    m = re.search(r"ARTWORK\s*REF:?\s*(.+?\.\s*AI)\b", t, re.I)
    if m:
        art = re.sub(r"\s+", "", m.group(1))
        refs.append("Art " + art)
    for m in re.finditer(r"DRAWING\s+(\d{5})", t, re.I):
        refs.append(f"Dwg {m.group(1)}")
    for m in re.finditer(r"\b(SDD-\d+)\b", t, re.I):
        refs.append(m.group(1).upper())

    size = ""
    m = re.search(r"(\d{2,4})\s*(?:mm)?\s*[xX\u00d7]\s*(\d{2,4})\s*mm", t)
    if m:
        size = f"{m.group(1)} x {m.group(2)}mm"
    elif art and re.match(r"(\d{2,4})x(\d{2,4})[-_]", art, re.I):
        a = re.match(r"(\d{2,4})x(\d{2,4})", art, re.I)
        size = f"{a.group(1)} x {a.group(2)} (artwork ref)"
    else:
        m = re.search(r"(?<![+\d])(\d{2,4})\s*mm\s+HIGH", t, re.I)
        if m:
            size = f"{m.group(1)}mm high"
        else:
            m = re.search(r"(?<![+\d.])(\d{2,4})\s*mm\b(?!\s*(?:AFFL|AFL|FFL|THICK|BACK|FROM|TO U/S|OFF))", t, re.I)
            if m and m.group(1) not in ("150",) or (m and "HUNG" not in t[max(0, m.start() - 5):m.start()].upper()):
                size = f"{m.group(1)}mm"
    ffl = ""
    m = re.search(r"(\d{3,4})\s*mm\s*(AFFL|AFL|FFL)", t, re.I)
    if m:
        ffl = f"{m.group(1)}mm {m.group(2).upper()}"
    else:
        m = re.search(r"\+\s*(\d{3,4})\s*mm\s*TO U/S", t, re.I)
        if m:
            ffl = f"+{m.group(1)}mm to underside"
    m = re.search(r"\(?\b(\d{1,3})\s*NO\b\.?\)?", t, re.I)
    stated = int(m.group(1)) if m else None

    flags = []
    if re.search(r"\bTB[CS]\b|TO BE CONFIRMED|SITE SIZED|SIZE TO BE|HEIGHT OF LETTER", t, re.I):
        flags.append("size/details TBC")
    if re.search(r"BY M&S|PROVIDED BY M&S|M&S MARKETING|ARTWORK DETAILS", t, re.I):
        flags.append("artwork by M&S")
    if re.search(r"TO BE ADDED|ON HOLD|AWAITING|CONFIRM LOCATIONS", t, re.I):
        flags.append("not yet defined")
    existing = bool(re.search(r"\bEXISTING\b|TO REMAIN|RE-?USE|RETAIN", t, re.I)) and not re.search(r"\bNEW\b", t, re.I)
    return {"refs": refs, "size": size, "ffl": ffl, "stated": stated, "flags": flags, "existing": existing}


def ms_tokens(t):
    return {w for w in re.findall(r"[A-Z0-9']{3,}", t.upper()) if w not in MS_STOP and not w.isdigit()}


MS_ADJ = set("ILLUMINATED BESPOKE HIGH WALL MOUNTED NEW HANGING NON NON-ILLUMINATED MID FLOOR FIXED FRAME AND THE A OF TO BE ON".split()) - {"ON"}


def ms_name(t):
    """Short name of a sign: the 1-3 words before SIGN/SIGNAGE (e.g. TRY ON, PAY HERE, LIFT)."""
    toks = re.findall(r"[A-Z&']+", re.sub(r"\d+\s*MM|\d+", " ", t.upper()))
    toks = [w.strip("'") for w in toks]
    for i, w in enumerate(toks):
        if w in ("SIGN", "SIGNAGE") and i:
            name = []
            for x in reversed(toks[max(0, i - 4):i]):
                if x in MS_ADJ or x in ("HIGH",):
                    if name:
                        break
                    continue
                name.append(x)
            if name:
                return " ".join(reversed(name))
    return ""


def ms_norm(t):
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z0-9 ]", "", t.upper())).strip()


def _ms_scan(pdf_bytes):
    """Read every page: title, drawing no., level, kind and sign callouts (text only - fast)."""
    doc = pdfium.PdfDocument(pdf_bytes)
    pages, items, nums_seen = [], [], Counter()
    seen_sigs = set()      # a pack bound twice repeats whole sheets - repeats are skipped
    for i in range(1, len(doc) + 1):
        pg = doc[i - 1]
        try:
            runs, dw, dh = page_runs(pg)
            text = " ".join(r["t"] for r in runs)
            title, num = ms_title(runs, dw, dh)
            sig = hashlib.md5(" ".join(sorted(r["t"] for r in runs)).encode()).hexdigest() if len(runs) > 5 else None
            dup = sig is not None and (title, sig) in seen_sigs
            if sig is not None:
                seen_sigs.add((title, sig))
            kind = "duplicate" if dup else ms_page_kind(title, text)
            level = ms_level(title, text)
            n = 0
            if kind in ("plan", "elevation"):
                for c in ms_cluster(runs):
                    t = c["text"]
                    if len(t) < 5 or c["x0"] > 0.83 * dw or not MS_SIGN_RE.search(t):
                        continue
                    if MS_NOISE_RE.search(t) or MS_LEGEND_RE.match(t.strip()) or t.strip().lower() == title.lower():
                        continue
                    if c["x0"] < 0.13 * dw and len(t) < 60 and re.search(r"KEY|EXAMPLE|NOTE", t, re.I):
                        continue
                    if len(t) > 420:
                        continue
                    n += 1
                    items.append({"page": i, "kind": kind, "level": level, "dwg": num, "title": title,
                                  "text": t, "box": (c["x0"], c["y0"], c["x1"], c["y1"]), "dw": dw, "dh": dh,
                                  "view": nearest_view_title([r for r in runs if "elevation" in r["t"].lower()],
                                                             (c["x0"], c["y0"], c["x1"], c["y1"])) if kind == "elevation" else ""})
            pages.append({"page": i, "type": kind, "level": level, "dwg": num, "title": title[:70], "callouts": n})
            nums_seen.update(num[:1] if isinstance(num, list) else [])
        except Exception as e:
            pages.append({"page": i, "type": "error", "level": "", "dwg": "", "title": str(e)[:60], "callouts": 0})
        finally:
            pg.close()
    doc.close()
    # the consultant project number repeats on every sheet - the drawing number is the other 5-digit value
    proj = nums_seen.most_common(1)[0][0] if nums_seen else ""
    for p in pages:
        if isinstance(p["dwg"], list):
            p["dwg"] = next((n for n in p["dwg"] if n != proj), "")
    for it in items:
        d = next((p["dwg"] for p in pages if p["page"] == it["page"]), "")
        it["dwg"] = d
    return {"pages": pages, "items": items}


ms_scan = st.cache_data(show_spinner="Reading binder...")(_ms_scan)



def _analyse_ms(pdf_bytes, include_pages, wpct, hpct, incl_plan, incl_elev, drop=()):
    sc = ms_scan(pdf_bytes)
    pages = sc["pages"]
    auto = {p["page"] for p in pages if p["type"] in ("plan", "elevation") and p["callouts"]}
    sel = set(include_pages) if include_pages else auto
    items = [it for it in sc["items"] if it["page"] in sel]
    if drop:   # callouts already accounted for (signage key box, artwork refs consumed by a numbered circle)
        items = [it for it in items if not any(it["page"] == p and _ms_rect_hit(it["box"], bx) for p, bx in drop)]

    # ---- group identical callouts (same sign called out on several plans / floors)
    groups = []
    for it in sorted(items, key=lambda x: (x["page"], x["box"][1], x["box"][0])):
        it["norm"] = ms_norm(it["text"])
        g = next((g for g in groups if g["kind_class"] == (it["kind"] == "plan")
                  and (g["norm"] == it["norm"] or difflib.SequenceMatcher(None, g["norm"], it["norm"]).ratio() >= 0.9)), None)
        if g is None:
            g = {"norm": it["norm"], "text": it["text"], "kind_class": it["kind"] == "plan", "items": []}
            groups.append(g)
        g["items"].append(it)

    plan_groups = [g for g in groups if g["kind_class"]]
    elev_groups = [g for g in groups if not g["kind_class"]]

    # ---- link elevation callouts to plan signs by shared distinctive words
    for g in plan_groups:
        tk = ms_tokens(g["text"])
        best, best_s = None, 0.0
        nm = ms_name(g["text"])
        for e in elev_groups:
            te = ms_tokens(e["text"])
            if not tk or not te:
                continue
            sh = len(tk & te)
            sc_ = sh / min(len(tk), len(te))
            lv_ok0 = (not e["items"][0]["level"])
            if nm and ms_name(e["text"]) == nm and (lv_ok0 or not g["items"][0]["level"] or e["items"][0]["level"] == g["items"][0]["level"]
                                                       or True):
                best, best_s = e, 9.0 if best_s < 9.0 else best_s
                continue
            lv_ok = (not e["items"][0]["level"]) or (not g["items"][0]["level"]) or e["items"][0]["level"] == g["items"][0]["level"]
            if lv_ok and sh >= 2 and sc_ >= 0.6 and sc_ > best_s:
                best, best_s = e, sc_
        g["elev_link"] = best
        if best is not None:
            best["linked"] = True
    elev_only = [e for e in elev_groups if not e.get("linked")]

    # ---- rows
    def rank(g):
        it = g["items"][0]
        return ({"B": 0, "GF": 1, "FF": 2}.get(it["level"], 3), it["page"], it["box"][1], it["box"][0])

    seq = Counter()
    rows, snips = [], {}
    doc = pdfium.PdfDocument(pdf_bytes)
    page_cache = {}

    def getpage(n):
        if n not in page_cache:
            page_cache[n] = doc[n - 1]
        return page_cache[n]

    def snippet(it, kind):
        pg = getpage(it["page"])
        hw, hh = wpct * it["dw"], hpct * it["dw"]
        scale = max(0.4, min(3.0, 1000.0 / (2 * hw)))
        cx, cy = (it["box"][0] + it["box"][2]) / 2, (it["box"][1] + it["box"][3]) / 2
        return {"kind": kind, "page": it["page"], "sheet": it["dwg"], "view": it["view"] or it["title"][:40],
                "label": f"{kind} p{it['page']} {it['dwg']} {it['view'] or it['title'][:40]}".strip(),
                "png": crop_png(pg, (cx, cy), hw, hh, scale, [it["box"]])}

    for g in sorted(plan_groups + elev_only, key=rank):
        its = g["items"]
        p = ms_parse(g["text"])
        if g["kind_class"] and g.get("elev_link"):
            pe = ms_parse(g["elev_link"]["text"])
            for k in ("size", "ffl"):
                if not p[k] and pe[k]:
                    p[k] = pe[k]
            if p["size"] and pe["size"] and re.findall(r"\d+", p["size"])[:1] != re.findall(r"\d+", pe["size"])[:1]:
                p["flags"] = p["flags"] + [f"size differs: plan {p['size']} / elevation {pe['size']}"]
            p["refs"] += [r for r in pe["refs"] if r not in p["refs"]]
        levels = []
        for it in its:
            if it["level"] and it["level"] not in levels:
                levels.append(it["level"])
        levels = sorted(levels, key=lambda l: {"B": 0, "GF": 1, "FF": 2}.get(l, 3))
        prefix = (levels[0] if levels else "EL") if g["kind_class"] else "EL"
        seq[prefix] += 1
        code = f"{prefix}-{seq[prefix]:02d}"
        n_plan = len(its) if g["kind_class"] else 0
        per_page = Counter(it["page"] for it in its)
        counted = n_plan if g["kind_class"] else max(per_page.values())
        qty = p["stated"] if (p["stated"] and counted == 1) else counted
        per_level = Counter(it["level"] or "?" for it in its)
        pos_bits = []
        if g["kind_class"]:
            pos_bits.append(", ".join(f"{LEVEL_NAMES.get(l, l)} x{n}" for l, n in per_level.items()))
        if g.get("elev_link"):
            pos_bits.append("elev: " + g["elev_link"]["items"][0]["view"] if g["elev_link"]["items"][0]["view"] else "elevation")
        if not g["kind_class"]:
            pos_bits.append(its[0]["view"] or its[0]["title"])
        issues = []
        if not p["size"]:
            issues.append("no size on callout")
        if p["existing"]:
            issues.append("existing / retained")
        issues += p["flags"]
        if g["kind_class"] and not g.get("elev_link"):
            issues.append("no matching elevation callout")
        if not g["kind_class"]:
            issues.append("elevation only - not on a signage plan")
        pages_txt = sorted({(it["page"], it["kind"], it["dwg"]) for it in its})
        if g.get("elev_link"):
            pages_txt += sorted({(it["page"], "elevation", it["dwg"]) for it in g["elev_link"]["items"]})
        r = {"level": ", ".join(levels) if levels else "-", "code": code, "desc": sentence_case(g["text"]),
             "qty": int(qty), "mod": "; ".join(p["refs"]), "sched_page": its[0]["page"],
             "size_name": p["size"], "ffl_name": p["ffl"], "position": "; ".join(pos_bits),
             "plan_check": (f"{n_plan} callout(s) on plans" if g["kind_class"] else f"{counted} callout(s) on elevations"),
             "check": "; ".join(issues),
             "pages": ", ".join(f"p{a} {b} {c}".strip() for a, b, c in pages_txt)}
        rows.append(r)

        lst = []
        if incl_plan and g["kind_class"]:
            lst.append(snippet(its[0], "plan"))
        if incl_elev:
            src = g["elev_link"]["items"][0] if g.get("elev_link") else (its[0] if not g["kind_class"] else None)
            if src is not None:
                lst.append(snippet(src, "elevation"))
        snips[f"{r['level']}|{code}"] = lst
        r["key"] = f"{r['level']}|{code}"

    for pg in page_cache.values():
        pg.close()
    doc.close()
    warnings = []
    if not rows:
        warnings.append("No sign callouts found on the selected pages.")
    return {"rows": rows, "snips": snips, "unmatched": [], "notes": [], "warnings": warnings,
            "pages": pages, "schedule_pages": [], "fallback": False, "hits": [], "selected": sorted(sel)}


analyse_ms = st.cache_data(show_spinner="Building M&S schedule...")(_analyse_ms)


# --------------------------------------------------------------------------------------
# M&S v2: signage KEY driven (numbered circles on the elevations -> numbered key entries)
# --------------------------------------------------------------------------------------
GENERIC_KEY_RE = re.compile(r"see elevations|graphics|artwork|local brand|lightbox", re.I)


def join_line_runs(rs):
    """Join text runs on one line; runs split mid-word by the PDF ('opaq' + 'ue') are glued back together."""
    rs = sorted(rs, key=lambda r: r["x0"])
    out = ""
    for k, r in enumerate(rs):
        if k:
            gap = r["x0"] - rs[k - 1]["x1"]
            out += (" " if gap > 0.30 * max(r["h"], 1.0) else "") + r["t"]
        else:
            out = r["t"]
    return out


def lines_from_runs(rs):
    rs = sorted(rs, key=lambda r: (r["y0"] + r["y1"]) / 2)
    lines, cur, cy = [], [], None
    for r in rs:
        c = (r["y0"] + r["y1"]) / 2
        if cur and abs(c - cy) <= 0.6 * r["h"]:
            cur.append(r)
        else:
            if cur:
                lines.append(cur)
            cur, cy = [r], c
    if cur:
        lines.append(cur)
    return [join_line_runs(l) for l in lines]


def find_bubbles(page, runs, dmin=8.0, dmax=40.0):
    """Numbers (1-2 digits) enclosed in a near-square/round vector path -> list of bubble dicts."""
    tf, dw, dh = geom(page)
    sq = []
    for o in page.get_objects(max_depth=3):
        if o.type == 2:                                   # path
            l, b, r, t = o.get_bounds()
            x0, y0, x1, y1 = box_tf(tf, l, b, r, t)
            w, h = x1 - x0, y1 - y0
            if dmin < w < dmax and h > 0 and 0.85 < w / h < 1.18:
                sq.append((x0, y0, x1, y1))
    out = []
    for r in runs:
        if not re.fullmatch(r"\d{1,2}", r["t"]):
            continue
        cx, cy = (r["x0"] + r["x1"]) / 2, (r["y0"] + r["y1"]) / 2
        enc = [s for s in sq if s[0] < cx < s[2] and s[1] < cy < s[3]]
        if enc:
            s = min(enc, key=lambda s: s[2] - s[0])
            out.append({"n": int(r["t"]), "cx": (s[0] + s[2]) / 2, "cy": (s[1] + s[3]) / 2,
                        "d": s[2] - s[0], "th": r["h"], "box": s})
    return out


def find_loose_numbers(runs, bubbles, valid, th, rect=None):
    """Numbers that match a key number and the key's label text size but whose circle outline wasn't recognised
    (some sign circles are drawn with curve segments the shape filter rejects)."""
    out = []
    for r in runs:
        if not re.fullmatch(r"\d{1,2}", r["t"]) or int(r["t"]) not in valid or abs(r["h"] - th) > 0.6:
            continue
        cx, cy = (r["x0"] + r["x1"]) / 2, (r["y0"] + r["y1"]) / 2
        if any(abs(b["cx"] - cx) < 9 and abs(b["cy"] - cy) < 9 for b in bubbles):
            continue
        box = (cx - 7, cy - 7, cx + 7, cy + 7)
        if rect and _ms_rect_hit(box, rect):
            continue
        out.append({"n": int(r["t"]), "cx": cx, "cy": cy, "d": 14.2, "th": r["h"], "box": box, "loose": True})
    return out


def ms_parse_key(runs, bubbles):
    """Read a 'SIGNAGE KEY' box: numbered entries '(n) Name / spec'. Returns (entries, rect) or ([], None)."""
    heads = [r for r in runs if re.fullmatch(r"SIGNAGE KEY", r["t"].strip(), re.I)]
    if not heads:
        return [], None
    hd = heads[0]
    kb = sorted([b for b in bubbles if b["cy"] > hd["y1"] - 2 and abs(b["cx"] - hd["x0"]) < 40], key=lambda b: b["cy"])
    seq, want = [], 1
    for b in kb:                                          # a key counts up 1, 2, 3 ... - stop at the first break
        if b["n"] == want:
            seq.append(b)
            want += 1
    if not seq:
        return [], None
    xlim = hd["x0"] + 260
    tidy = lambda x: re.sub(r"(?<=[a-z]) ?- ?(?=[a-z])", "-", x)
    entries = []
    for i, b in enumerate(seq):
        y0 = b["cy"] - b["d"] / 2 - 3
        y1 = seq[i + 1]["cy"] - seq[i + 1]["d"] / 2 - 3 if i + 1 < len(seq) else b["cy"] + 75
        rs = [r for r in runs if b["cx"] + b["d"] / 2 - 2 < r["x0"] < xlim and y0 <= (r["y0"] + r["y1"]) / 2 < y1]
        ls = [l for l in lines_from_runs(rs) if l.strip(" -")]
        if not ls:
            continue
        entries.append({"n": b["n"], "name": tidy(ls[0].strip()), "desc": tidy(" ".join(ls[1:]).strip()),
                        "text": tidy(" ".join(ls).strip()), "bd": b["d"], "th": b["th"]})
    rect = (hd["x0"] - 12, hd["y0"] - 8, xlim, seq[-1]["cy"] + 80)
    return entries, rect


def ms_view_label(t):
    t = re.sub(r"^(?:[A-Z]\s+)+(?=PROPOSED)", "", t.strip())
    t = re.sub(r"^PROPOSED\s+", "", t)
    return " ".join(w if re.fullmatch(r"[A-Z]|\(.*\)", w) else w.capitalize() for w in t.split())


def ms_view_titles(runs):
    """View titles ('PROPOSED INTERNAL ELEVATION C') - left-anchored, set just under their view."""
    return [r for r in runs if len(r["t"]) < 70 and r["h"] > 9 and re.search(r"\bPROPOSED\b", r["t"])
            and re.search(r"\b(?:ELEVATION|SECTION)\b", r["t"]) and not re.search(r"SIGNAGE|DETAILS", r["t"])]


def ms_view_title(runs, pt, titles=None):
    """The view a point sits in: nearest title below it; in that row the right-most title to its left."""
    titles = titles if titles is not None else ms_view_titles(runs)
    cx, cy = pt
    c = [(r["y0"] - cy, r) for r in titles if r["y0"] >= cy - 5 and r["x0"] <= cx + 40]
    if not c:
        return ""
    m = min(d for d, _ in c)
    row = [r for d, r in c if d <= m + 60]
    return ms_view_label(max(row, key=lambda r: r["x0"])["t"])


def _ms_rect_hit(box, rect, pad=0.0):
    return not (box[2] < rect[0] - pad or box[0] > rect[2] + pad or box[3] < rect[1] - pad or box[1] > rect[3] + pad)


def ms_size_from_text(t):
    m = re.search(r"(\d{3,4})\s*[xX×]\s*(\d{3,4})\s*mm", t)
    if m:
        return f"{m.group(1)} x {m.group(2)}mm"
    m = re.search(r"\b(\d{2,4})\s*mm\b", t)
    return f"{m.group(1)}mm" if m else ""


def _ms_keyscan(pdf_bytes):
    """Which pages carry a SIGNAGE KEY, and what does it say?"""
    doc = pdfium.PdfDocument(pdf_bytes)
    out = {}
    try:
        for i in range(1, len(doc) + 1):
            pg = doc[i - 1]
            try:
                runs, dw, dh = page_runs(pg)
                if not any(re.fullmatch(r"SIGNAGE KEY", r["t"].strip(), re.I) for r in runs):
                    continue
                entries, rect = ms_parse_key(runs, find_bubbles(pg, runs))
                if entries:
                    out[i] = {"entries": entries, "rect": rect}
            finally:
                pg.close()
    finally:
        doc.close()
    return out


ms_keyscan = st.cache_data(show_spinner="Looking for the signage key...")(_ms_keyscan)


def _analyse_ms_key(pdf_bytes, key_page, apply_pages, wpct, hpct, incl_plan, incl_elev, extra_pages):
    """Key-driven M&S schedule. key_page = sheet holding the master SIGNAGE KEY; apply_pages = sheets whose numbered
    circles refer to it; extra_pages = sheets also read for text callouts that are not in the key."""
    ks = ms_keyscan(pdf_bytes)
    entries, key_rect = ks[key_page]["entries"], ks[key_page]["rect"]
    by_n = {e["n"]: e for e in entries}
    kd_main = sorted(e["bd"] for e in entries)[len(entries) // 2]
    scan = ms_scan(pdf_bytes)
    page_info = {p["page"]: p for p in scan["pages"]}
    doc = pdfium.PdfDocument(pdf_bytes)
    hits, drop, notes = [], [], []
    try:
        for pn in apply_pages:
            pg = doc[pn - 1]
            try:
                runs, dw, dh = page_runs(pg)
                own = ks.get(pn)                       # a sheet with its own key box: use it for the circle size
                kd = sorted(e["bd"] for e in own["entries"])[len(own["entries"]) // 2] if own else kd_main
                allb = find_bubbles(pg, runs)
                bub = [b for b in allb if abs(b["d"] - kd) <= 0.12 * kd and b["n"] in by_n]
                kth = sorted(e["th"] for e in entries)[len(entries) // 2]
                bub += find_loose_numbers(runs, allb, set(by_n), kth, (own or ks[key_page])["rect"])
                if own:
                    bub = [b for b in bub if not _ms_rect_hit(b["box"], own["rect"])]
                    drop.append((pn, own["rect"]))
                    for e2 in own["entries"]:
                        m = by_n.get(e2["n"])
                        if m and ms_norm(m["name"]) != ms_norm(e2["name"]):
                            notes.append(f"Key on p{pn} disagrees with key on p{key_page} at item {e2['n']}: "
                                         f"'{e2['name']}' vs '{m['name']}'")
                arts = [r for r in runs if re.search(r"ARTWORK\s*REF", r["t"], re.I)]
                titles = ms_view_titles(runs)
                pairs = []                             # artwork ref under a bubble whose entry is 'see elevations for graphics'
                for bi, b in enumerate(bub):
                    if not GENERIC_KEY_RE.search(by_n[b["n"]]["desc"]):
                        continue
                    for ai, a in enumerate(arts):
                        dx = abs((a["x0"] + a["x1"]) / 2 - b["cx"])
                        dy = a["y0"] - b["cy"]
                        if dx < 220 and 0 < dy < 650:
                            pairs.append((dx + 0.15 * dy, bi, ai))
                assigned, taken = {}, set()
                for _, bi, ai in sorted(pairs):
                    if bi not in assigned and ai not in taken:
                        assigned[bi] = ai
                        taken.add(ai)
                for bi, b in enumerate(bub):
                    art = ""
                    if bi in assigned:
                        a = arts[assigned[bi]]
                        art = re.sub(r"(?i)^artwork\s*ref:?\s*", "", a["t"]).strip()
                        art = re.sub(r"\s+", "", art) if "_" in art else art
                        drop.append((pn, (a["x0"] - 4, a["y0"] - 30, a["x1"] + 4, a["y1"] + 14)))
                    hits.append({"loose": b.get("loose", False), "n": b["n"], "page": pn, "cx": b["cx"], "cy": b["cy"], "d": b["d"], "art": art,
                                 "view": ms_view_title(runs, (b["cx"], b["cy"]), titles), "dw": dw, "dh": dh,
                                 "box": b["box"], "dwg": page_info.get(pn, {}).get("dwg", ""),
                                 "title": page_info.get(pn, {}).get("title", "")})
            finally:
                pg.close()
    finally:
        doc.close()
    drop.append((key_page, key_rect))
    notes = list(dict.fromkeys(notes))

    rows, snips = [], {}
    doc = pdfium.PdfDocument(pdf_bytes)
    pcache = {}

    def getpage(n):
        if n not in pcache:
            pcache[n] = doc[n - 1]
        return pcache[n]

    def snippet(h):
        pg = getpage(h["page"])
        hw, hh = wpct * h["dw"], hpct * h["dw"]
        scale = max(0.4, min(3.0, 1000.0 / (2 * hw)))
        return {"kind": "elevation", "page": h["page"], "sheet": h["dwg"], "view": h["view"] or h["title"][:40],
                "label": f"elevation p{h['page']} {h['dwg']} {h['view']}".strip(),
                "png": crop_png(pg, (h["cx"], h["cy"]), hw, hh, scale, [h["box"]])}

    def position(hs):
        c = Counter(h["view"] or f"p{h['page']}" for h in hs)
        return "; ".join(f"{v} x{n}" if n > 1 else v for v, n in c.items())

    for e in entries:
        hs = sorted([h for h in hits if h["n"] == e["n"]], key=lambda h: (h["page"], h["cy"], h["cx"]))
        base_issues = []
        if re.search(r"\bTB[CS]\b|TO BE CONFIRMED|SITE SIZED|SIZE TO BE", e["text"], re.I):
            base_issues.append("size/details TBC")
        if re.search(r"artwork supplied by|by marketing|by M&S", e["text"], re.I):
            base_issues.append("artwork by M&S")
        groups = [(None, hs)]
        if hs and GENERIC_KEY_RE.search(e["desc"]) and any(h["art"] for h in hs):
            g = defaultdict(list)
            for h in hs:
                g[h["art"] or ""].append(h)
            groups = list(g.items())
        for gi, (art, ghs) in enumerate(groups, 1):
            code = f"K{e['n']:02d}" + (f".{gi}" if len(groups) > 1 else "")
            issues = list(base_issues)
            if not hs:
                issues.append("in the signage key but no numbered circle found on the selected sheets")
            if art == "" and len(groups) > 1:
                issues.append("artwork ref not found next to this sign")
            if any(h["loose"] for h in ghs):
                issues.append(f"{sum(1 for h in ghs if h['loose'])} circle(s) read from the number only - check")
            per_sheet = Counter(h["page"] for h in ghs)
            qty = max(per_sheet.values()) if per_sheet else 0
            if len(per_sheet) > 1 and len(set(per_sheet.values())) > 1:
                issues.append("count differs between sheets (" + ", ".join(f"p{p}: {n}" for p, n in sorted(per_sheet.items()))
                              + ") - check")
            r = {"level": "-", "code": code, "desc": e["name"] + (f" - {art}" if art else ""), "qty": qty,
                 "mod": f"Key {e['n']}" + (f"; Art {art}" if art else ""),
                 "sched_page": ghs[0]["page"] if ghs else key_page,
                 "size_name": ms_size_from_text(e["desc"]), "ffl_name": "",
                 "position": position(ghs) if ghs else "",
                 "plan_check": (f"{len(ghs)} numbered circle(s) on " + ", ".join(f"p{p}" for p in sorted(per_sheet)))
                 if ghs else "none found",
                 "check": "; ".join(issues),
                 "pages": ", ".join(sorted({f"p{h['page']} {h['dwg']}".strip() for h in ghs})) + f"; key p{key_page}",
                 "spec": e["desc"], "key": f"-|{code}"}
            rows.append(r)
            snips[r["key"]] = [snippet(ghs[0])] if (ghs and incl_elev) else []
    for pg in pcache.values():
        pg.close()
    doc.close()

    # text callouts that are not part of the key (external elevations, door artwork ...)
    if extra_pages:
        base = analyse_ms(pdf_bytes, tuple(sorted(extra_pages)), wpct, hpct, incl_plan, incl_elev,
                          tuple((p, tuple(b)) for p, b in drop))
        keep_re = re.compile(r"ARTWORK\s*REF|ILLUMINATED|\bLED\b|VINYL|ACRYLIC|TRAY|DIAMETER|\d+\s*MM|SUPPLIER", re.I)
        kept = [r for r in base["rows"] if keep_re.search(r["desc"])]
        for r in kept:
            r["check"] = ("text callout, not in signage key; " + r["check"]).strip("; ")
        rows += kept
        snips.update({k: v for k, v in base["snips"].items() if k in {r["key"] for r in kept}})

    warnings = []
    missing = [e for e in entries if not any(h["n"] == e["n"] for h in hits)]
    if missing:
        warnings.append("Key items with no numbered circle on the selected sheets: "
                        + ", ".join(f"{e['n']} {e['name']}" for e in missing[:12]) + (" ..." if len(missing) > 12 else ""))
    return {"rows": rows, "snips": snips, "unmatched": [], "notes": notes, "warnings": warnings,
            "pages": scan["pages"], "schedule_pages": [key_page], "fallback": False, "hits": hits,
            "selected": sorted(set(apply_pages) | set(extra_pages)), "key_entries": entries}


analyse_ms_key = st.cache_data(show_spinner="Building M&S schedule from the signage key...")(_analyse_ms_key)


# --------------------------------------------------------------------------------------
# M&S visual sweep (optional, needs an API key): reads rendered tiles of the elevations and reports every
# key sign it can SEE - including signs drawn as artwork with no number and no callout (e.g. a FLOWER SHOP
# sign that is just white lettering on a wall). The vector pass above only finds what is numbered or labelled.
# --------------------------------------------------------------------------------------
SWEEP_PROMPT = """This is a crop (tile) of a store signage elevation drawing. Below is the SIGNAGE KEY for the store.

{key}

Report every sign from this key that is physically drawn in THIS tile, whether or not it has a numbered circle beside it.
Match by what the sign says or looks like (its lettering, shape, illumination) and by its numbered circle when it has one.
Also report anything that clearly looks like a sign (lettering, graphic panel, roundel) but matches nothing in the key.
Do NOT report: dimensions, grid-line circles, finish tags like WF:04, shelving, products, wall or floor finishes, red clouds.
Return JSON only, no other text:
{{"signs":[{{"key_no":<int or null>,"seen_text":"<wording on the sign>","circle_number":<int or null>,
"x_pct":<0-100>,"y_pct":<0-100>,"confidence":"high"|"medium"|"low"}}]}}
Use key_no null for signs not in the key. x_pct / y_pct are the centre of the sign within this tile (0,0 = top-left).
Only report what is clearly visible. If you are unsure whether something is a sign, leave it out."""


def sweep_tiles(region, tw=1100.0, th=780.0, ov=110.0):
    x0, y0, x1, y1 = region
    xs, ys = [], []
    x = x0
    while x < x1 - 40:
        xs.append(x)
        x += tw - ov
    y = y0
    while y < y1 - 40:
        ys.append(y)
        y += th - ov
    return [(a, b, min(a + tw, x1), min(b + th, y1)) for b in ys for a in xs]


def _sweep_read(client, model, jpg: bytes, key_text: str):
    msg = client.messages.create(
        model=model, max_tokens=1800,
        messages=[{"role": "user", "content": [
            {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": base64.b64encode(jpg).decode()}},
            {"type": "text", "text": SWEEP_PROMPT.format(key=key_text)}]}])
    text = "".join(b.text for b in msg.content if getattr(b, "type", "") == "text")
    a, b = text.find("{"), text.rfind("}")
    if a < 0 or b < 0:
        return []
    try:
        return json.loads(text[a:b + 1]).get("signs", [])
    except Exception:
        return []


def ms_visual_sweep(client, model, pdf_bytes, page_no, entries, region, progress=None, scale=1.6, reader=None):
    """{"seen": [...], "unmatched": [...], "tiles": n}. `reader(jpg, key_text)` can be injected for testing."""
    key_text = "\n".join(f"{e['n']}: {e['name']} - {e['desc'][:110]}" for e in entries)
    doc = pdfium.PdfDocument(pdf_bytes)
    seen, unmatched = [], []
    tiles = sweep_tiles(region)
    try:
        pg = doc[page_no - 1]
        for ti, box in enumerate(tiles, 1):
            im = crop_image(pg, box, scale)
            buf = io.BytesIO()
            im.save(buf, format="JPEG", quality=82)
            try:
                signs = (reader or (lambda j, k: _sweep_read(client, model, j, k)))(buf.getvalue(), key_text)
            except Exception as e:
                signs = []
                unmatched.append({"page": page_no, "key_no": None, "text": f"tile {ti} failed: {e}",
                                  "x": box[0], "y": box[1], "conf": "", "circle": None})
            tw_, th_ = box[2] - box[0], box[3] - box[1]
            for s_ in signs:
                try:
                    x = box[0] + float(s_.get("x_pct", 50)) / 100 * tw_
                    y = box[1] + float(s_.get("y_pct", 50)) / 100 * th_
                except Exception:
                    continue
                rec = {"page": page_no, "key_no": s_.get("key_no"), "text": str(s_.get("seen_text") or ""),
                       "x": x, "y": y, "conf": s_.get("confidence", ""), "circle": s_.get("circle_number")}
                (seen if isinstance(rec["key_no"], int) else unmatched).append(rec)
            if progress:
                progress(ti / len(tiles))
        pg.close()
    finally:
        doc.close()
    out = []                                       # overlapping tiles see the same sign twice - merge close detections
    for r in sorted(seen, key=lambda r: (r["key_no"], r["y"], r["x"])):
        if not any(o["key_no"] == r["key_no"] and math.hypot(o["x"] - r["x"], o["y"] - r["y"]) < 170 for o in out):
            out.append(r)
    uo = []
    for r in unmatched:
        if not any(math.hypot(o["x"] - r["x"], o["y"] - r["y"]) < 170 and o["text"].lower() == r["text"].lower() for o in uo):
            uo.append(r)
    return {"seen": out, "unmatched": uo, "tiles": len(tiles)}


def merge_sweep(res, sweep, pdf_bytes, wpct, hpct):
    """Fold visual-sweep results into a copy of the key-driven result. Never overwrites circle counts."""
    res = {**res, "rows": [dict(r) for r in res["rows"]], "snips": dict(res["snips"]), "warnings": list(res["warnings"])}
    seen = defaultdict(list)
    for r in sweep["seen"]:
        seen[r["key_no"]].append(r)
    doc = pdfium.PdfDocument(pdf_bytes)
    titles = {}
    try:
        for r in res["rows"]:
            m = re.fullmatch(r"K(\d+)", r["code"])
            if not m:
                continue
            vs = seen.get(int(m.group(1)), [])
            if not vs:
                continue
            if r["qty"] == 0:
                r["qty"] = len(vs)
                bits = []
                for v in vs:
                    if v["page"] not in titles:
                        pg = doc[v["page"] - 1]
                        titles[v["page"]] = ms_view_titles(page_runs(pg)[0])
                        pg.close()
                    bits.append(ms_view_title(None, (v["x"], v["y"]), titles[v["page"]]) or f"p{v['page']}")
                r["position"] = "; ".join(f"{k} x{c}" if c > 1 else k for k, c in Counter(bits).items())
                r["plan_check"] = "visual sweep only"
                r["check"] = ("seen by visual sweep only - no numbered circle or callout on the drawing, confirm; "
                              + r["check"]).strip("; ")
                r["mod"] = r["mod"] + "; visual"
                pg = doc[vs[0]["page"] - 1]
                tf, dw, dh = geom(pg)
                hw, hh = wpct * dw, hpct * dw
                scale = max(0.4, min(3.0, 1000.0 / (2 * hw)))
                box = (vs[0]["x"] - 14, vs[0]["y"] - 14, vs[0]["x"] + 14, vs[0]["y"] + 14)
                res["snips"][r["key"]] = [{"kind": "elevation", "page": vs[0]["page"], "sheet": "", "view": "",
                                           "label": f"visual sweep p{vs[0]['page']}",
                                           "png": crop_png(pg, (vs[0]["x"], vs[0]["y"]), hw, hh, scale, [box])}]
                pg.close()
            elif len(vs) > r["qty"]:
                r["check"] = (f"visual sweep sees {len(vs)}, circles give {r['qty']} - check; " + r["check"]).strip("; ")
    finally:
        doc.close()
    if sweep["unmatched"]:
        res["warnings"].append(f"Visual sweep: {len(sweep['unmatched'])} sign-like item(s) that match nothing in the key.")
    res["sweep_unmatched"] = sweep["unmatched"]
    return res


# --------------------------------------------------------------------------------------
# Optional: Vision read
# --------------------------------------------------------------------------------------
VISION_PROMPT = """This is a crop from a store signage elevation drawing. The red box marks the tag "{code}" ({desc}).
Read ONLY what is drawn or written on the drawing. Do not infer or assume.
Return JSON only, with these keys:
  "position": short phrase - which wall/area/view this sign is on
  "size": this sign's size as dimensioned (with units), or null
  "ffl": its height from finished floor level as dimensioned, and what it is measured to, or null
  "tbc": any "TBC", "by others" or similar note attached to it, or null
  "confidence": "high", "medium" or "low"
  "evidence": the exact dimension text you used, or null
Use null for anything not shown. General drawing notes:
{notes}"""


def vision_read(client, model, png: bytes, row: dict, notes: list):
    msg = client.messages.create(
        model=model, max_tokens=700,
        messages=[{"role": "user", "content": [
            {"type": "image", "source": {"type": "base64", "media_type": "image/png",
                                       "data": base64.b64encode(png).decode()}},
            {"type": "text", "text": VISION_PROMPT.format(
                code=row["code"], desc=row["desc"], notes="\n".join(notes[:6]) or "(none)")},
        ]}])
    text = "".join(b.text for b in msg.content if getattr(b, "type", "") == "text")
    a, b = text.find("{"), text.rfind("}")
    return json.loads(text[a:b + 1])


COLS = ["Level", "Sign Code", "Description", "Size (text)", "Size (drawing)", "Position",
        "FFL", "Qty", "Plan check", "Check", "MOD / Ref", "Read confidence", "Evidence"]


def build_df(res, vision):
    hang_note = next((n for n in res["notes"] if "hanging" in n.lower()), "")
    out = []
    for r in res["rows"]:
        v = vision.get(f"{r['level']}|{r['code']}", {})
        ffl = v.get("ffl") or r.get("ffl_name", "")
        if not ffl and "hanging" in r["desc"].lower() and hang_note:
            ffl = f"per note: {hang_note}"
        if v.get("tbc"):
            ffl = f"{ffl} [{v['tbc']}]".strip()
        out.append({
            "Level": r["level"], "Sign Code": r["code"], "Description": r["desc"],
            "Size (text)": r["size_name"], "Size (drawing)": v.get("size") or "",
            "Position": r["position"], "FFL": ffl, "Qty": r["qty"], "Plan check": r["plan_check"],
            "Check": r["check"], "MOD / Ref": r["mod"], "Read confidence": v.get("confidence") or "",
            "Evidence": v.get("evidence") or "",
        })
    return pd.DataFrame(out, columns=COLS)


def build_xlsx(df, snips) -> bytes:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Signage Schedule"
    headers = COLS + ["Snippet"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    widths = [7, 12, 46, 20, 18, 34, 28, 6, 18, 24, 8, 12, 24, 48]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    keep = []
    for n, (_, row) in enumerate(df.iterrows(), start=2):
        ws.append([row[c] for c in COLS])
        for c in ws[n]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        key = f"{row['Level']}|{row['Sign Code']}"
        if snips.get(key):
            im = Image.open(io.BytesIO(snips[key][0]["png"]))
            im.thumbnail((360, 260))
            b = io.BytesIO()
            im.save(b, format="PNG")
            b.seek(0)
            xi = XLImage(b)
            keep.append(b)
            ws.add_image(xi, f"{ws.cell(row=1, column=len(headers)).column_letter}{n}")
            ws.row_dimensions[n].height = 200
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_zip(snips) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for key, lst in snips.items():
            level, code = key.split("|")
            for i, s in enumerate(lst, 1):
                z.writestr(f"{level}_{code}_{s['kind']}_p{s['page']}_{i}.png", s["png"])
    return out.getvalue()


def build_pdf(df, snips, source_name="", summary=True) -> bytes:
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import (Image as RLImage, PageBreak, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], fontSize=20, leading=24, alignment=0, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=ss["Normal"], fontSize=12, leading=15, spaceAfter=6)
    cell = ParagraphStyle("cell", parent=ss["Normal"], fontSize=7.5, leading=9)
    cellb = ParagraphStyle("cellb", parent=cell, fontName="Helvetica-Bold")
    fact = ParagraphStyle("fact", parent=ss["Normal"], fontSize=9.5, leading=12)
    cap = ParagraphStyle("cap", parent=ss["Normal"], fontSize=8, leading=10, textColor=colors.grey)

    def P(text, style=cell):
        return Paragraph(escape(str(text or "")), style)

    page_w, page_h = landscape(A4)
    margin = 28
    frame_w = page_w - 2 * margin
    img_area_h = page_h - 2 * margin - 135

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.grey)
        canvas.drawString(margin, 14, f"Signage schedule - {source_name}")
        canvas.drawRightString(page_w - margin, 14, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=margin, rightMargin=margin,
                            topMargin=margin, bottomMargin=margin, title="Signage schedule")
    story = []

    if summary:
        story.append(Paragraph("Signage Schedule", h1))
        story.append(Paragraph(escape(source_name), h2))
        head = ["Lvl", "Code", "Description", "Size", "FFL", "Qty", "Position", "Check"]
        data = [[P(h, cellb) for h in head]]
        for _, r in df.iterrows():
            data.append([P(r["Level"]), P(r["Sign Code"], cellb), P(r["Description"]),
                         P(r["Size (drawing)"] or r["Size (text)"]), P(r["FFL"]), P(r["Qty"]),
                         P(r["Position"]), P(r["Check"] or r["Plan check"])])
        t = Table(data, colWidths=[28, 60, 190, 100, 112, 26, 150, 120], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDDDDD")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F5")]),
        ]))
        story.append(t)
        story.append(PageBreak())

    def fit(png, max_w, max_h):
        w, h = Image.open(io.BytesIO(png)).size
        k = min(max_w / w, max_h / h)
        return RLImage(io.BytesIO(png), width=w * k, height=h * k)

    n_rows = len(df)
    for n, (_, r) in enumerate(df.iterrows(), start=1):
        story.append(Paragraph(f"{escape(r['Sign Code'])} &nbsp;&nbsp;<font size=11 color='#555555'>"
                               f"{escape(r['Level'])} &nbsp;|&nbsp; Qty {r['Qty']}</font>", h1))
        story.append(Paragraph(escape(r["Description"]), h2))
        size = r["Size (drawing)"] or r["Size (text)"] or "-"
        def F(label, value):
            return Paragraph(f"<b>{label}:</b> {escape(str(value))}", fact)
        facts = Table([[F("Size", size), F("FFL", r["FFL"] or "-")],
                       [F("Position", r["Position"] or "-"), F("Check", r["Check"] or r["Plan check"])]],
                      colWidths=[frame_w * 0.5, frame_w * 0.5])
        facts.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                   ("BOTTOMPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 0)]))
        story.append(facts)
        story.append(Spacer(1, 8))

        lst = snips.get(f"{r['Level']}|{r['Sign Code']}", [])[:2]
        if not lst:
            story.append(Paragraph("No plan or elevation reference found for this sign.", fact))
        else:
            each_w = frame_w / len(lst) - 6
            cells = [[[fit(s["png"], each_w, img_area_h - 14), Paragraph(escape(s["label"]), cap)]
                     for s in lst]]
            it = Table(cells, colWidths=[frame_w / len(lst)] * len(lst))
            it.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("ALIGN", (0, 0), (-1, -1), "CENTER")]))
            story.append(it)
        if n < n_rows:
            story.append(PageBreak())
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# --------------------------------------------------------------------------------------
# UI
# --------------------------------------------------------------------------------------
def get_api_key() -> str:
    try:
        if "ANTHROPIC_API_KEY" in st.secrets:
            return st.secrets["ANTHROPIC_API_KEY"]
    except Exception:
        pass
    return os.environ.get("ANTHROPIC_API_KEY", "")


def run_vision(res, keys, api_key, model, state):
    import anthropic
    client = anthropic.Anthropic(api_key=api_key)
    by_key = {f"{r['level']}|{r['code']}": r for r in res["rows"]}
    bar = st.progress(0.0, text="Reading drawings...")
    for n, key in enumerate(keys, 1):
        snips = res["snips"].get(key, [])
        snips = [s for s in snips if s["kind"] == "elevation"] or snips
        if snips and key not in state:
            try:
                state[key] = vision_read(client, model, snips[0]["png"], by_key[key], res["notes"])
            except Exception as e:
                state[key] = {"evidence": f"read failed: {e}"}
        bar.progress(n / len(keys), text=f"Reading drawings... {n}/{len(keys)}")
    bar.empty()


def vision_sidebar():
    st.divider()
    st.subheader("Vision read (optional)")
    api_key = get_api_key() or st.text_input("Anthropic API key", type="password")
    model = st.selectbox("Model", ["claude-haiku-5-5", "claude-sonnet-5-5"])
    return api_key, model


def show_results(res, up_name, file_id, mode, api_key, model):
    for w in res.get("warnings", []):
        st.warning(w)
    if not res["rows"]:
        st.error("Nothing found. Check you've uploaded the right pack for this mode.")
        return

    vstore = st.session_state.setdefault("vision", {})
    vision = vstore.setdefault(f"{mode}:{file_id}", {})
    df_all = build_df(res, vision)
    df_all["_key"] = [r.get("key") or f"{r['level']}|{r['code']}" for r in res["rows"]]

    dropped = st.session_state.setdefault("dropped", {}).setdefault(f"{mode}:{file_id}", [])
    df = df_all[~df_all["Sign Code"].isin(dropped)].reset_index(drop=True)
    snips = {k: v for k, v in res["snips"].items() if k in set(df["_key"])}
    df = df.drop(columns=["_key"])
    res_view = {**res, "rows": [r for r in res["rows"] if r["code"] not in dropped], "snips": snips}

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Signs", len(df))
    c2.metric("Total qty", int(df["Qty"].sum()))
    c3.metric("With snippets", sum(1 for v in snips.values() if v))
    c4.metric("Flagged", int((df["Check"] != "").sum()))

    t_sched, t_signs, t_check = st.tabs(["Schedule", "Sign by sign", "Check sheet"])

    with t_sched:
        st.dataframe(df, use_container_width=True, hide_index=True)
        new_drop = st.multiselect("Remove rows that aren't real signs (removed from every export)",
                                  list(df_all["Sign Code"]), default=dropped,
                                  format_func=lambda c: f"{c}  -  {df_all.loc[df_all['Sign Code'] == c, 'Description'].iloc[0][:70]}")
        if new_drop != dropped:
            st.session_state["dropped"][f"{mode}:{file_id}"] = new_drop
            st.rerun()
        d1, d2, d3, d4 = st.columns(4)
        d1.download_button("Excel (with snippets)", build_xlsx(df, snips), "signage_schedule.xlsx",
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        d2.download_button("PDF (one page per sign)", build_pdf(df, snips, up_name), "signage_schedule.pdf",
                           "application/pdf")
        d3.download_button("Snippets (zip)", build_zip(snips), "snippets.zip", "application/zip")
        d4.download_button("CSV", df.to_csv(index=False).encode("utf-8"), "signage_schedule.csv", "text/csv")
        st.divider()
        if api_key:
            if st.button("Read sizes / FFL from the drawing snippets (all signs)", key=f"va_{mode}"):
                run_vision(res_view, list(snips.keys()), api_key, model, vision)
                st.rerun()
            st.caption("Vision reads only what's drawn on the snippet - always check against the evidence column.")
        else:
            st.caption("Add an Anthropic API key in the sidebar to read sizes and FFL off the snippets.")

    with t_signs:
        if df.empty:
            st.info("No signs left.")
        else:
            labels = [f"{r['Level']}  {r['Sign Code']}  -  {r['Description'][:90]}" for _, r in df.iterrows()]
            idx_key = f"sign_idx_{mode}"
            st.session_state.setdefault(idx_key, 0)
            st.session_state[idx_key] = min(st.session_state[idx_key], len(labels) - 1)

            def step(d):
                st.session_state[idx_key] = (st.session_state[idx_key] + d) % len(labels)

            n1, n2, n3 = st.columns([1, 6, 1])
            n1.button("< Prev", on_click=step, args=(-1,), use_container_width=True, key=f"prev_{mode}")
            n3.button("Next >", on_click=step, args=(1,), use_container_width=True, key=f"next_{mode}")
            n2.selectbox("Sign", range(len(labels)), format_func=lambda i: labels[i], key=idx_key,
                         label_visibility="collapsed")
            i = st.session_state[idx_key]
            r = df.iloc[i]
            key = f"{r['Level']}|{r['Sign Code']}"
            st.subheader(f"{r['Sign Code']}  |  {r['Level']}  |  Qty {r['Qty']}")
            st.write(r["Description"])
            f1, f2 = st.columns(2)
            f1.markdown(f"**Size:** {r['Size (drawing)'] or r['Size (text)'] or '-'}")
            f1.markdown(f"**FFL:** {r['FFL'] or '-'}")
            f1.markdown(f"**Ref:** {r['MOD / Ref'] or '-'}")
            f2.markdown(f"**Position:** {r['Position'] or '-'}")
            f2.markdown(f"**Check:** {r['Check'] or r['Plan check']}")
            if r["Evidence"]:
                st.caption(f"Evidence: {r['Evidence']} ({r['Read confidence']} confidence)")
            if api_key and st.button("Read this sign from its snippet", key=f"vr_{mode}_{key}"):
                run_vision(res_view, [key], api_key, model, vision)
                st.rerun()
            lst = snips.get(key, [])
            if not lst:
                st.warning("No plan or elevation reference found for this sign.")
            else:
                for col, s in zip(st.columns(len(lst)), lst):
                    col.image(s["png"], caption=s["label"], use_container_width=True)

    with t_check:
        flagged = df[df["Check"] != ""]
        if flagged.empty:
            st.success("Nothing flagged.")
        else:
            st.markdown("**Signs to look at**")
            st.dataframe(flagged[["Level", "Sign Code", "Description", "Qty", "Plan check", "Check"]],
                         use_container_width=True, hide_index=True)
        if res["unmatched"]:
            st.markdown("**Schedule lines that looked like signs but didn't parse**")
            st.code("\n".join(res["unmatched"]))
        if res["notes"]:
            st.markdown("**FFL notes found on the drawings**")
            for n in res["notes"][:15]:
                st.write(f"- {n}")
        with st.expander("How each page was classified"):
            st.dataframe(pd.DataFrame(res["pages"]), use_container_width=True, hide_index=True)


def tk_mode():
    with st.sidebar:
        st.header("Settings")
        elev_n = st.slider("Elevation snippets per sign", 1, 3, 1)
        incl_plan = st.checkbox("Also include a plan snippet", value=True)
        ehw = st.slider("Elevation crop half-width (pt)", 50, 600, 220, 10)
        ehh = st.slider("Elevation crop half-height (pt)", 50, 600, 160, 10)
        phw = st.slider("Plan crop half-width (pt)", 50, 600, 150, 10)
        phh = st.slider("Plan crop half-height (pt)", 50, 600, 110, 10)
        dpi = st.select_slider("Snippet resolution (dpi)", options=[100, 150, 200, 250], value=150)
        api_key, model = vision_sidebar()

    up = st.file_uploader("Upload the TK Maxx drawing pack (PDF)", type=["pdf"], key="tk_pdf")
    if not up:
        st.info("Upload a pack to get the schedule, one page per sign, and a check sheet.")
        return
    pdf_bytes = up.getvalue()
    file_id = hashlib.md5(pdf_bytes).hexdigest()
    try:
        res = analyse(pdf_bytes, elev_n, incl_plan, ehw, ehh, phw, phh, dpi)
    except Exception as e:
        st.error(f"Couldn't read that PDF: {e}")
        return
    if not res["rows"]:
        st.error("No schedule rows found. This pack has no 'TK_Signage' schedule (Level / Family and Type / TK Code / Count) - "
                 "some packs, such as the Beauty refit packs, don't include one. Check it's the right pack.")
        return
    if res["fallback"]:
        st.warning("No schedule page recognised - scanned every page for rows instead. Check results carefully.")
    show_results(res, up.name, file_id, "tk", api_key, model)


def ms_mode():
    with st.sidebar:
        st.header("Settings")
        incl_plan = st.checkbox("Plan snippet (where it is)", value=True)
        incl_elev = st.checkbox("Elevation snippet (where one matches)", value=True)
        wpct = st.slider("Snippet half-width (% of sheet width)", 3.0, 20.0, 9.0, 0.5) / 100
        hpct = st.slider("Snippet half-height (% of sheet width)", 2.0, 15.0, 6.5, 0.5) / 100
        api_key, model = vision_sidebar()

    up = st.file_uploader("Upload the M&S binder (PDF)", type=["pdf"], key="ms_pdf")
    if not up:
        st.info("Upload the M&S binder. Signs are picked up from the red callouts on the signage plans and elevations, "
                "counted across floors, with a plan and elevation snippet for each.")
        return
    pdf_bytes = up.getvalue()
    file_id = hashlib.md5(pdf_bytes).hexdigest()
    try:
        scan = ms_scan(pdf_bytes)
    except Exception as e:
        st.error(f"Couldn't read that PDF: {e}")
        return

    pages = scan["pages"]
    auto = [p["page"] for p in pages if p["type"] in ("plan", "elevation") and p["callouts"]]
    ndup = sum(1 for p in pages if p["type"] == "duplicate")
    if ndup:
        st.info(f"{ndup} page(s) repeat an earlier sheet exactly (the pack looks bound twice) and are ignored, so nothing is counted twice.")
    by_page = {p["page"]: p for p in pages}
    with st.expander(f"Pages being read ({len(auto)} auto-selected) - change if a sign sheet was missed or a wrong one included"):
        sel = st.multiselect(
            "Pages", [p["page"] for p in pages if p["type"] != "error"], default=auto,
            format_func=lambda n: f"p{n} | {by_page[n]['dwg'] or '-'} | {by_page[n]['type']} | {by_page[n]['title'] or '(no title found)'} | {by_page[n]['callouts']} callouts")
    if not sel:
        st.warning("Pick at least one page.")
        return
    ks = ms_keyscan(pdf_bytes)
    use_key = False
    if ks:
        key_page = max(ks, key=lambda p: len(ks[p]["entries"]))
        use_key = st.checkbox(f"Use the signage key on p{key_page} ({len(ks[key_page]['entries'])} sign types) - recommended",
                              value=True, help="Every sign in the key gets a row, found by its numbered circle on the elevations.")
    try:
        if use_key:
            with st.expander("Sheets whose numbered circles refer to the key"):
                apply_pages = st.multiselect(
                    "Sheets", [p["page"] for p in pages if p["type"] != "error"], default=sorted(ks), key="ms_apply",
                    format_func=lambda n: f"p{n} | {by_page[n]['dwg'] or '-'} | {by_page[n]['title'] or '(no title found)'}")
            if not apply_pages:
                st.warning("Pick at least one sheet.")
                return
            res = analyse_ms_key(pdf_bytes, key_page, tuple(sorted(apply_pages)), wpct, hpct, incl_plan, incl_elev,
                                 tuple(sorted(sel)))
            if api_key:
                with st.expander("Visual sweep - finds signs that have no number or label on the drawing", expanded=False):
                    sw_pages = st.multiselect("Sheets to sweep", apply_pages, default=[key_page], key="ms_sweep_pages",
                                              format_func=lambda n: f"p{n} {by_page[n]['title'][:50]}")
                    skey = f"{file_id}|{model}|{tuple(sw_pages)}"
                    store = st.session_state.setdefault("sweeps", {})
                    st.caption("Reads the drawing as an image, tile by tile, and lists every key sign it can see. "
                               "Anything it adds is flagged 'visual sweep only' - always check the snippet.")
                    if st.button("Run visual sweep", key="ms_sweep_go"):
                        import anthropic
                        client = anthropic.Anthropic(api_key=api_key)
                        bar = st.progress(0.0, text="Sweeping the drawings...")
                        acc = {"seen": [], "unmatched": [], "tiles": 0}
                        d_ = pdfium.PdfDocument(pdf_bytes)
                        for k_, pn in enumerate(sw_pages):
                            pg_ = d_[pn - 1]
                            _, dw_, dh_ = geom(pg_)
                            pg_.close()
                            part = ms_visual_sweep(client, model, pdf_bytes, pn, ks[key_page]["entries"],
                                                   (0, 0, 0.66 * dw_, 0.96 * dh_),
                                                   progress=lambda f, k_=k_: bar.progress((k_ + f) / len(sw_pages)))
                            for kk in ("seen", "unmatched"):
                                acc[kk] += part[kk]
                            acc["tiles"] += part["tiles"]
                        d_.close()
                        bar.empty()
                        store[skey] = acc
                        st.rerun()
                    if skey in store:
                        res = merge_sweep(res, store[skey], pdf_bytes, wpct, hpct)
                        st.success(f"Sweep done: {len(store[skey]['seen'])} key signs seen across {store[skey]['tiles']} tiles.")
                        if res.get("sweep_unmatched"):
                            st.markdown("**Sign-like items that match nothing in the key**")
                            st.dataframe(pd.DataFrame(res["sweep_unmatched"])[["page", "text", "conf"]],
                                         use_container_width=True, hide_index=True)
            else:
                st.caption("Add an Anthropic API key in the sidebar to switch on the visual sweep for signs with no number or label.")
        else:
            res = analyse_ms(pdf_bytes, tuple(sorted(sel)), wpct, hpct, incl_plan, incl_elev)
    except Exception as e:
        st.error(f"Couldn't build the M&S schedule: {e}")
        return
    if use_key:
        st.caption("Sign IDs here are the M&S signage key numbers (K03 = key item 3). Quantity is the number of numbered circles "
                   "on the elevations (the highest count on any one sheet, so a sign drawn on two sheets isn't doubled). "
                   "Anything flagged in the Check column is worth a look.")
    else:
        st.caption("M&S doesn't use sign codes, so each sign gets an ID (B / GF / FF + number; EL = only seen on an elevation). "
                   "The M&S UPC / artwork / drawing reference is in the 'MOD / Ref' column. Quantity is the number of callouts "
                   "counted on the signage plans across all floors unless a callout states a number.")
    show_results(res, up.name, file_id, "ms", api_key, model)


# --------------------------------------------------------------------------------------
# Manual Checker + Revision Checker
# --------------------------------------------------------------------------------------
# ---- Manual reading (Geetee Signs sheets) ----------------------------------------------------
GT_DWG_RE = re.compile(r"\b[A-Z]{2}-\d{4}-\d{2,4}\b")
GT_QTY_RE = re.compile(r"(\d+)\s*(?:no\.?|off|nr\.?)\s*(?:full\s+)?(?:(?:sets?|each|pairs?|panels?)\s*)*(?:required|req\b\.?|total)", re.I)
GT_LIN_RE = re.compile(r"(\d+(?:\.\d+)?)\s*linear\s*m", re.I)
GT_FLAG_RE = re.compile(r"\b(ON HOLD|TBC|TO BE CONFIRMED|ARTWORK (?:REQUIRED|TBC|TO FOLLOW)|AWAITING|NOT REQUIRED|REMOVED|DELETED|OMITTED)\b", re.I)
GT_LAYOUT_RE = re.compile(r"elevation|key plan|plan\b|^sign [a-z]$|manifestation|opening hours|legal signage|\bcover\b|index|contents", re.I)


def gt_join_runs(runs):
    """Rejoin text pieces the PDF split mid-word ('W'+'indow', 'Pegboar'+'d'): same line, similar size, touching (gap < 0.2 x height)."""
    out = []
    for r in sorted(runs, key=lambda r: r["x0"]):
        best, bg = None, None
        for p in out:
            g = r["x0"] - p["x1"]
            h = max(p["h"], r["h"])
            if (-0.5 <= g < 0.2 * h and abs((p["y0"] + p["y1"]) / 2 - (r["y0"] + r["y1"]) / 2) < 0.35 * h
                    and 0.7 < r["h"] / p["h"] < 1.43 and (bg is None or abs(g) < abs(bg))):
                best, bg = p, g
        if best is not None:
            best["t"] = best["t"] + r["t"]
            best["x1"] = max(best["x1"], r["x1"]); best["y0"] = min(best["y0"], r["y0"]); best["y1"] = max(best["y1"], r["y1"])
            best["h"] = max(best["y1"] - best["y0"], 1.0)
        else:
            out.append(dict(r))
    return out


def gt_char_text(page, d):
    """Title and body text in the PDF's own character order (cleaner than text boxes, which split words and repeat letters).
    Returns (title, body). Title = big text right of DESCRIPTION:, body = everything outside the title-block strip."""
    try:
        import pypdfium2.raw as pdfium_c
    except Exception:
        return "", ""
    tf, dw, dh = geom(page)
    tp = page.get_textpage()
    title, body = [], []
    try:
        for i in range(tp.count_chars()):
            c = tp.get_text_range(i, 1)
            if not c:
                continue
            if c in "\r\n":
                (title if title and title[-1] != " " else []).append(" ")
                if body and body[-1] != " | ":
                    body.append(" | ")
                continue
            l, b, r, t = tp.get_charbox(i)
            x0, y0, x1, y1 = box_tf(tf, l, b, r, t)
            if y0 > d["y0"] - 62 and x0 > d["x0"] - 17:
                if d["x0"] + 3 <= x0 < d["x0"] + 183 and d["y0"] - 4 <= y0 <= d["y0"] + 26 and pdfium_c.FPDFText_GetFontSize(tp.raw, i) > 7:
                    title.append(c)
                continue
            body.append(c)
    finally:
        tp.close()
    t = re.sub(r"\s+", " ", "".join(title)).strip()
    t = re.split(r"\s*This drawing", t)[0].strip()
    bt = "".join(body).replace("\ufffe", "-").replace("\x02", "-")
    bt = " | ".join(x.strip() for x in re.split(r"\s*\|\s*", bt) if x.strip() and not x.strip().startswith(("This drawing", "All measurements")))
    return t, re.sub(r"[ \t]+", " ", bt)


def read_manual_page(page, pno):
    runs, dw, dh = page_runs(page)
    runs = gt_join_runs(runs)
    d = next((r for r in runs if r["t"].upper().startswith("DESCRIPTION:")), None)
    out = {"page": pno, "dwg": "", "title": "", "qty": [], "lin": [], "flags": [], "revs": [], "text": "", "nums": [], "ffl": [], "is_sheet": d is not None}
    if d is None:
        out["text"] = " ".join(r["t"] for r in runs)[:400]
        return out
    # title: runs right of the DESCRIPTION label, in the same block row
    dx = d["x0"] - 497.5                                  # title blocks are not always in the same place on the sheet
    tl = [r for r in runs if 500 + dx <= r["x0"] < 680 + dx and d["y0"] - 4 <= r["y0"] <= d["y0"] + 26 and r["h"] > 7
          and not r["t"].upper().endswith(":") and not re.fullmatch(r"\d\d\.\d\d\.\d{2,4}|n/a", r["t"].strip(), re.I)
          and "geetee" not in r["t"].lower()]
    tl.sort(key=lambda r: (round(r["y0"] / 3), r["x0"]))
    t = " ".join(r["t"] for r in tl)
    t = re.split(r"\s*This drawing", t)[0].strip()
    ct, cbody = gt_char_text(page, d)
    if re.search(r"[A-Za-z]{3,}", ct):
        t = ct
    out["title"] = t
    m = [r for r in runs if GT_DWG_RE.fullmatch(r["t"].strip())]
    if m:
        out["dwg"] = m[0]["t"].strip()
    # revision rows (rev table sits right of x=850 in the title block)
    rt = [r for r in runs if 850 + dx <= r["x0"] <= 1000 + dx and r["y0"] > d["y0"] - 60]
    dates = [r for r in rt if re.fullmatch(r"\d\d\.\d\d\.\d\d", r["t"]) and not r["t"].startswith("00.00")]
    for dt in sorted(dates, key=lambda r: r["y0"]):
        row = [r for r in rt if abs(r["y0"] - dt["y0"]) < 4 and r["x0"] > dt["x1"]]
        row.sort(key=lambda r: r["x0"])
        letter = next((r["t"] for r in rt if abs(r["y0"] - dt["y0"]) < 6 and r["x0"] < dt["x0"] and len(r["t"]) == 1), "")
        txt = [r["t"] for r in row if not re.fullmatch(r"X+", r["t"])]
        by = txt[0] if txt and len(txt[0]) <= 3 else ""
        desc = " ".join(txt[1:] if by else txt)
        out["revs"].append({"rev": letter, "date": dt["t"], "by": by, "desc": desc})
    # body = everything outside the title block strip
    body = [r for r in runs if not (r["y0"] > d["y0"] - 62 and r["x0"] > 480 + dx) and r["t"] != t]
    body = [r for r in body if not r["t"].startswith("This drawing") and not r["t"].startswith("All measurements")]
    txt = " | ".join(r["t"] for r in sorted(body, key=lambda r: (round(r["y0"] / 6), r["x0"])))
    if cbody:
        txt = cbody
    out["text"] = txt
    qtxt = re.sub(r"\s*\|\s*(?:\.\s*)?", " ", txt)          # text runs split mid-sentence ("1 No | . required")
    for mm in GT_QTY_RE.finditer(qtxt):
        out["qty"].append(int(mm.group(1)))
    out["lin"] = [float(x) for x in GT_LIN_RE.findall(txt)]
    out["ffl"] = sorted({int(x) for x in re.findall(r"(\d{3,4})\s*mm\s*(?:AFFL|FFL|A\.F\.F\.L)", qtxt, re.I)} |
                        {int(x) for x in re.findall(r"(?:AFFL|FFL)[^\d|]{0,12}(\d{3,4})\s*mm", qtxt, re.I)})
    out["flags"] = sorted({f.upper() for f in GT_FLAG_RE.findall(txt)})
    for r in body:
        for n in re.findall(r"(?<![\d.])(\d{2,5})(?:\s*mm)?(?![\d.])", r["t"]):
            out["nums"].append(int(n))
    return out


CK_VER = "2026-10-14f"   # bump when the readers change: Streamlit only hashes a cached function's own body, so stale results survive edits to helpers


def _read_manual(pdf_bytes, ver=""):
    pdf = pdfium.PdfDocument(pdf_bytes)
    pages = []
    try:
        for i in range(len(pdf)):
            pg = pdf[i]
            try:
                pages.append(read_manual_page(pg, i + 1))
            finally:
                pg.close()
    finally:
        pdf.close()
    return pages


read_manual = st.cache_data(show_spinner="Reading your manual...")(_read_manual)


def manual_signs(pages):
    """Group sheets into signs by (drawing no, title). Layout-only sheets are kept apart."""
    signs, layout = {}, []
    for p in pages:
        if not p["is_sheet"] or not re.search(r"[A-Za-z]{3,}", p["title"] or ""):
            continue
        spec = bool(re.search(r"scale 1:|specification|vinyl|letters|required|fixed|mm", p["text"], re.I))
        has_q = bool(p["qty"] or p["lin"])
        if GT_LAYOUT_RE.search(p["title"]) and not has_q:
            layout.append(p)
            continue
        key = (p["dwg"], p["title"].lower())
        s = signs.setdefault(key, {"title": p["title"], "dwg": p["dwg"], "pages": [], "qty": [], "lin": [], "flags": [],
                                   "revs": [], "text": "", "nums": [], "ffl": [], "spec": False})
        s["pages"].append(p["page"])
        s["qty"] += p["qty"]; s["lin"] += p["lin"]; s["flags"] = sorted(set(s["flags"]) | set(p["flags"]))
        s["ffl"] = sorted(set(s["ffl"]) | set(p.get("ffl", []))); s["revs"] += p["revs"]; s["text"] += " | " + p["text"]; s["nums"] += p["nums"]; s["spec"] |= spec
    out = [s for s in signs.values() if s["spec"] or s["qty"]]
    layout += [p for s in signs.values() if not (s["spec"] or s["qty"]) for p in []]
    return out, layout


# ---- Matching manual signs to pack rows -------------------------------------------------------
_CK_STOP = set("a an the and or of to for in on at with by from is are be as no mm m x sign signs signage vinyl vinyls graphic graphics "
               "illuminated non wall mounted hanging fixed required req sets set each off new existing see refer artwork detail details "
               "applied internal internally external supply install fit fitted type size high wide long centres floor level ground first "
               "one two three four five six tbc edit text message ref uk eu nav dpt".split())
_CK_SYN = {"toilets": "toilet", "wc": "toilet", "wcs": "toilet", "loo": "toilet", "bakers": "baker", "bakery": "baker", "baking": "bake",
           "tills": "till", "checkout": "till", "colleagues": "colleague", "customers": "customer", "doors": "door", "mirrors": "mirror",
           "lifts": "lift", "slats": "slat", "totems": "directory", "totem": "directory", "directory": "directory", "directories": "directory", "boh": "boh",
           "backofhouse": "boh", "pay": "pay", "payment": "pay", "fitting": "fitting", "rooms": "room", "windows": "window",
           "letters": "letter", "lettering": "letter", "manifestations": "manifestation", "roundel": "roundel", "roundels": "roundel"}


def _ck_stem(w):
    w = _CK_SYN.get(w, w)
    if len(w) > 4 and w.endswith("ing") and w not in _CK_SYN:
        w = w[:-3]
    elif len(w) > 4 and w.endswith("ed"):
        w = w[:-2]
    if len(w) > 3 and w.endswith("s") and w not in _CK_SYN:
        w = w[:-1]
    if len(w) > 4 and w.endswith("er") and w not in ("number", "water", "corner"):
        w = w[:-2]
    return w


_CK_STOP |= {_ck_stem(w) for w in set(_CK_STOP)} | set("bespoke affl front back bulkhead high hang hung colour color grey gray black white matt "
                                                       "gloss acrylic aluminium frame paint finish cut flat fret thick vinyls applied internal".split())


def ck_tokens(s, keep_generic=False):
    out = set()
    s = re.sub(r"\S+\.(?:jpe?g|png|ai|pdf|eps|tif+|svg)\b", " ", (s or ""), flags=re.I)
    for w in re.findall(r"[a-z0-9]+", s.lower().replace("&", " and ")):
        w = _ck_stem(w)
        if not keep_generic and (w in _CK_STOP or w.isdigit() or len(w) < 2 or (re.search(r"\d", w) and not re.fullmatch(r"[a-z]{1,3}\d+", w))):
            continue
        out.add(w)
    return out


def ck_idf(signs, rows):
    import math
    docs = [ck_tokens(s["title"]) for s in signs] + [ck_tokens(r["desc"]) for r in rows]
    n = max(len(docs), 1)
    df = {}
    for d in docs:
        for w in d:
            df[w] = df.get(w, 0) + 1
    return {w: math.log(1 + n / c) for w, c in df.items()}


def ck_score(sign, row, idf=None):
    """0..1 - how well a pack row's description fits a manual sign (idf-weighted word overlap)."""
    generic = not ck_tokens(sign["title"])
    if not ck_tokens(row["desc"]):
        return 0.0                                    # 'GF-05 : vinyl' - nothing distinctive to match a sign by
    tt = ck_tokens(sign["title"]) or ck_tokens(sign["title"], True)
    rt = ck_tokens(row["desc"]) or ck_tokens(row["desc"], True)
    if not tt or not rt:
        return 0.0
    rt = {next((w for w in tt if len(w) >= 6 and len(x) >= 6 and difflib.SequenceMatcher(None, w, x).ratio() >= 0.85), x) for x in rt}
    def _pairs(txt):
        ws = [_ck_stem(w) for w in re.findall(r"[a-z0-9]+", (txt or "").lower())]
        return {_ck_stem(a + b): (a, b) for a, b in zip(ws, ws[1:])}
    rp, tp_ = _pairs(row["desc"]), _pairs(sign["title"])
    for w in list(tt):
        if w in rp:                                   # sheet 'buzzbars' = pack 'buzz bars'
            rt = (rt - set(rp[w])) | {w}
    for w in list(rt):
        if w in tp_:                                  # sheet 'peg board' = pack 'pegboard'
            tt = (tt - set(tp_[w])) | {w}
    if generic and sign.get("text"):
        # Generic title ('Hanging signage'): the sign names are in the drawing. Match when the pack's distinctive words are all on the
        # sheet and the generic word ('hang') is shared.
        bt = ck_tokens(sign["text"])
        gw = ck_tokens(sign["title"], True) & ck_tokens(row["desc"], True) - {"sign", "signage"}
        if rt and rt <= bt and gw:
            return 0.85
        if rt and gw and len(rt & bt) / len(rt) >= 0.5:
            return 0.55
    wt = (lambda w: idf.get(w, 1.0)) if idf else (lambda w: 1.0)
    shared = tt & rt
    glued = re.sub(r"[^a-z0-9]", "", (sign["title"] or "").lower())    # titles with split letters ("W elcome", "Fr eestanding")
    more = {w for w in rt if len(w) >= 4 and w in glued} - shared
    if more:
        shared = shared | more
        tt = tt | more
    if not shared:
        return 0.0
    if len(rt) <= 3 and sign.get("text") and rt <= (tt | ck_tokens(sign["text"])) and len(rt & tt) >= 1:
        floor = 0.6                                   # short pack description entirely covered by the sheet's title + own text
    else:
        floor = 0.0
    base = sum(map(wt, shared)) / min(sum(map(wt, tt)), sum(map(wt, rt)))
    if base < 1.0 and sign.get("text"):
        bt = ck_tokens(sign["text"])
        extra = (rt - shared) & bt
        if extra:
            base = min(1.0, base + 0.1 * min(len(extra), 2))
    return round(min(max(base, floor), 1.0), 3)


def ck_match(signs, rows, thr=0.5):
    """Each pack row goes to its best-scoring sign(s) (ties kept). Returns {sign index: [rows]}, [unmatched rows]."""
    idf = ck_idf(signs, rows)
    got = {i: [] for i in range(len(signs))}
    left = []
    for r in rows:
        if r.get("qty", 0) <= 0:
            continue
        sc = [(ck_score(s, r, idf), i) for i, s in enumerate(signs)]
        best = max(sc, default=(0, -1))
        if best[0] >= thr:
            for v, i in sc:
                if v >= best[0] - 1e-9:
                    got[i].append(r)
        else:
            left.append(r)
    return got, left


# ---- Shared: read a pack into schedule rows ---------------------------------------------------
def _ck_pack(pdf_bytes, retailer, snippets, ver=""):
    if retailer == "TK Maxx":
        res = _analyse(pdf_bytes, 1, snippets, 220, 160, 150, 110, 100)
    else:
        scan = ms_scan(pdf_bytes)
        auto = tuple(sorted(p["page"] for p in scan["pages"] if p["type"] in ("plan", "elevation") and p["callouts"]))
        ks = ms_keyscan(pdf_bytes)
        if ks:
            key_page = max(ks, key=lambda p: len(ks[p]["entries"]))
            res = _analyse_ms_key(pdf_bytes, key_page, tuple(sorted(ks)), 0.09, 0.065, snippets, snippets, auto)
        else:
            res = _analyse_ms(pdf_bytes, auto, 0.09, 0.065, snippets, snippets)
    rows = []
    for r in res["rows"]:
        r = dict(r)
        r.setdefault("level", "-"); r.setdefault("position", ""); r.setdefault("check", ""); r.setdefault("size_name", "")
        r.setdefault("pages", ""); r.setdefault("desc", ""); r.setdefault("qty", 0); r.setdefault("code", "")
        r["key"] = r.get("key") or f"{r['level']}|{r['code']}"
        m = re.search(r"\bp(\d+)", r["pages"] or "")
        r["sched_page"] = int(m.group(1)) if m else r.get("sched_page", 0)
        rows.append(r)
    return {"rows": rows, "snips": res.get("snips", {}), "pages": res.get("pages", []), "warnings": res.get("warnings", []),
            "mode": "key" if (retailer != "TK Maxx" and res.get("key_entries")) else "callout"}


ck_pack = st.cache_data(show_spinner="Reading the pack...")(_ck_pack)


# ---- Manual Checker ---------------------------------------------------------------------------
def _ck_sizes(text):
    return {int(n) for n in re.findall(r"(?<![\d.])(\d{2,4})\s*mm", text or "", re.I)}


TK_CODE_RE = re.compile(r"TK\d{2}-\d{1,3}[A-Z]{0,2}(?:-[A-Z]{1,2})?|TKM-\d+[A-Z]*")


def _tk_base(c):
    m = re.match(r"(TK\d{2}-\d{1,3}|TKM-\d+)", c.upper())
    return m.group(1) if m else c.upper()


def tk_code_ok(pack_code, manual_code):
    """Same item: manual sheets write 'TK04-56A' where the pack writes 'TK04-56-A' (and may drop the colour suffix)."""
    return _tk_base(pack_code) == _tk_base(manual_code)


def tk_code_same(pack_code, manual_code):
    """Strict version for the side-by-side table: same item AND same variant letters (hyphens ignored), unless one side leaves the suffix off."""
    if not tk_code_ok(pack_code, manual_code):
        return False
    a, b = pack_code.upper()[len(_tk_base(pack_code)):].replace("-", ""), manual_code.upper()[len(_tk_base(manual_code)):].replace("-", "")
    return a == b or not a or not b


def ck_check(signs, rows, retailer):
    """Check each manual sign against the pack rows. Returns (checks, sign_summaries, pack_only)."""
    live = [r for r in rows if r.get("qty", 0) > 0]
    got, left = {i: [] for i in range(len(signs))}, []
    idf = ck_idf(signs, live)
    tk = retailer == "TK Maxx"
    sign_codes = [set(TK_CODE_RE.findall((x["title"] + " " + x["text"]).upper())) for x in signs] if tk else [set()] * len(signs)
    by_code = set()
    if tk:                                                     # pass 1: TK codes written on the manual sheets
        for r in live:
            hit = [i for i, cs in enumerate(sign_codes) if r["code"] and any(tk_code_ok(r["code"].upper(), m) for m in cs)]
            for i in hit:
                got[i].append((r, 1.0))
            if hit:
                by_code.add(id(r))
    code_found = [any(any(tk_code_ok(r["code"].upper(), m) for m in cs) for r in live) for cs in sign_codes]
    for r in live:                                             # pass 2: by name, for what the codes didn't settle
        if id(r) in by_code:
            continue
        sc = [(ck_score(s, r, idf), i) for i, s in enumerate(signs) if not (tk and sign_codes[i] and code_found[i])]
        best = max(sc, default=(0, -1))
        if best[0] >= 0.5:
            for v, i in sc:
                if v >= best[0] - 1e-9:
                    got[i].append((r, v))
        else:
            left.append(r)
    # pass 3: a sheet with a generic title ('Hanging signage') that found nothing takes the leftover pack rows of the same family
    # ('UK_Hanging Nav Sign_410 - Double line') as a loose match, so it shows once as 'confirm' instead of twice as missing
    for i, s in enumerate(signs):
        if got[i] or ck_tokens(s["title"]):
            continue
        gw = ck_tokens(s["title"], True) - {"sign", "signage"}
        fam = [r for r in left if gw & ck_tokens(r["desc"], True)]
        if gw and fam:
            got[i] = [(r, 0.55) for r in fam]
            left = [r for r in left if all(r is not f for f in fam)]
    checks, summ = [], []
    for i, s in enumerate(signs):
        rs = got[i]
        pk = [r if r.get("key") else {**r, "key": f"{r.get('level', '-')}|{r['code']}"} for r, _ in rs]
        conf = min((v for _, v in rs), default=0)
        mpg = ", ".join(f"p{p}" for p in s["pages"])
        name = s["title"] + (f"  [{s['dwg']}]" if s["dwg"] else "")
        ppg = ", ".join(sorted({f"p{r['sched_page']}" for r in pk if r["sched_page"]}, key=lambda x: int(x[1:])))
        base = {"Sign": s["title"], "Dwg": s["dwg"], "Manual page": mpg, "_keys": [r["key"] for r in pk if r.get("key")], "_mp": s["pages"][0]}
        rows_out = []
        if not pk:
            rows_out.append({**base, "Check": "In the pack", "Manual says": "Manual sheet " + mpg, "Pack says": "No matching callout / key item",
                             "Pack page": "", "Status": "Check",
                             "Note": "Sign is in your manual but not found in the pack - confirm it is still wanted, or the pack uses a different name"})
        else:
            names = "; ".join(f"{r['code']} {r['desc'][:45]}" for r in pk[:4]) + (f" (+{len(pk) - 4} more)" if len(pk) > 4 else "")
            rows_out.append({**base, "Check": "In the pack", "Manual says": s["title"], "Pack says": names, "Pack page": ppg,
                             "Status": "OK" if conf >= 0.7 else "Check",
                             "Note": "" if conf >= 0.7 else ("Your sheet has a general title - matched to the pack item(s) of the same type left over; confirm" if conf == 0.55 and not ck_tokens(s["title"]) else "Loose name match - confirm it is the same sign (quantity and size not compared)")})
            mq = sorted({q for q in s["qty"] if q})
            pq = sum(r["qty"] for r in pk)
            if conf < 0.7:
                mq, psz_skip = [], True
            if mq:
                ok = pq in mq or pq == sum(mq) or pq == sum(sorted(set(mq))) or any(pq == q * len(pk) for q in mq)
                rows_out.append({**base, "Check": "Quantity", "Manual says": " / ".join(f"{q} No." for q in mq), "Pack says": str(pq),
                                 "Pack page": ppg, "Status": "OK" if ok else "Check",
                                 "Note": "" if ok else "Quantities differ - check the sheet and the pack callouts"})
            elif s["lin"] and conf >= 0.7:
                ok = any(abs(pq - q) < 0.5 for q in s["lin"])
                rows_out.append({**base, "Check": "Quantity", "Manual says": " / ".join(f"{q:g} linear m" for q in s["lin"]),
                                 "Pack says": f"{pq}" + (" (schedule count)" if tk else " callout(s)"), "Pack page": ppg,
                                 "Status": "OK" if (ok or not tk) else "Check",
                                 "Note": "" if ok else ("Pack count differs from your linear metres - the pack may be counting panels, not metres; confirm" if tk else "Measured in linear metres - not compared")})
            psz = sorted({n for r in pk for n in _ck_sizes(("" if "(name)" in (r["size_name"] or "") else (r["size_name"] or "")) + " " + r["desc"])})
            if psz and conf >= 0.7:
                msz = set(s["nums"]) | _ck_sizes(s["text"])
                near = lambda a, b: abs(a - b) <= max(10, 0.03 * b)          # 748 on the sheet = '750mm high' in the pack
                miss = [n for n in psz if not any(near(m, n) for m in msz)]
                close = [n for n in miss if any(0.75 * n <= m <= 1.25 * n for m in msz)]
                if not miss:
                    st_, note = "OK", ""
                elif close:
                    st_, note = "Check", ", ".join(f"{n}mm" for n in close) + " in the pack - your sheet has " + \
                                ", ".join(str(m) for m in sorted(msz) if any(0.75 * n <= m <= 1.25 * n for n in close))
                else:
                    st_, note = "Info", ("Not compared - no matching dimension could be read as text on your sheet (dimensions drawn as graphics). "
                                         "Use the Drawing check below to read them from the drawing.")
                rows_out.append({**base, "Check": "Size", "Manual says": "sizes on sheet: " + (", ".join(str(n) for n in sorted(msz)[:8]) or "none read"),
                                 "Pack says": ", ".join(f"{n}mm" for n in psz), "Pack page": ppg, "Status": st_, "Note": note})
            if s.get("ffl"):
                pf = sorted({int(x) for r in pk for x in re.findall(r"(\d{3,4})\s*mm\s*(?:AFFL|FFL)", (r.get("ffl_name") or "") + " " + r["desc"], re.I)})
                if pf:
                    ok = any(f in s["ffl"] for f in pf)
                    rows_out.append({**base, "Check": "FFL height", "Manual says": ", ".join(f"{f}mm" for f in s["ffl"]),
                                     "Pack says": ", ".join(f"{f}mm" for f in pf), "Pack page": ppg, "Status": "OK" if ok else "Check",
                                     "Note": "" if ok else "Fixing height differs from the pack"})
                else:
                    rows_out.append({**base, "Check": "FFL height", "Manual says": ", ".join(f"{f}mm" for f in s["ffl"]),
                                     "Pack says": "no readable height in the pack text", "Pack page": ppg, "Status": "OK",
                                     "Note": "Not compared automatically - check against the pack drawing in the view below"})
            fl = []
            for r in pk:
                for kw in ("TBC", "not yet defined", "existing", "retained", "differs", "number only", "visual sweep", "artwork by M&S"):
                    if kw.lower() in r["check"].lower() and kw not in fl:
                        fl.append(kw)
            if fl:
                _fw = {"differs": "the pack's own schedule count differs from its plan callouts"}
                rows_out.append({**base, "Check": "Pack flags", "Manual says": "", "Pack says": "; ".join(_fw.get(f, f) for f in fl), "Pack page": ppg,
                                 "Status": "Check", "Note": "The pack itself flags this sign as unsettled or non-standard"})
        if tk and sign_codes[i]:
            absent = sorted(m for m in sign_codes[i] if not any(tk_code_ok(r["code"].upper(), m) for r in live))
            if absent:
                alt = [r for r in live if ck_score(s, r, idf) >= 0.6]
                fam = {m[:4] for m in absent}
                alt = alt or [r for r in left if r["code"][:4].upper() in fam]
                hint = ("; the pack lists a similar sign as " + ", ".join(sorted({r["code"] for r in alt})[:3])) if alt else ""
                rows_out.append({**base, "Check": "TK code", "Manual says": ", ".join(sorted(sign_codes[i])),
                                 "Pack says": "not in the schedule: " + ", ".join(absent) + hint, "Pack page": ppg, "Status": "Check",
                                 "Note": "Code on your sheet is not in the pack's schedule"})
        if s["flags"]:
            rows_out.append({**base, "Check": "Manual notes", "Manual says": ", ".join(s["flags"]), "Pack says": "", "Pack page": "",
                             "Status": "Check", "Note": "Your sheet carries a hold / TBC note"})
        checks += rows_out
        last = max(s["revs"], key=lambda r: (r["rev"] or ""), default=None) if s["revs"] else None
        summ.append({"sign": s["title"], "dwg": s["dwg"], "pages": mpg, "status": "Check" if any(r["Status"] == "Check" for r in rows_out) else "OK",
                     "n_pack": len(pk), "pk": pk, "rev": (f"Rev {last['rev']} {last['desc']}" if last else "")})
    pack_only = [{"Sign": r["desc"][:110], "Code": r["code"], "Qty": r["qty"], "Where": (r["position"] or r["level"])[:70],
                  "Pack page": f"p{r['sched_page']}" if r["sched_page"] else "", "Status": "Missing",
                  "Note": (r["check"] or "")[:110], "_key": r.get("key") or f"{r.get('level', '-')}|{r['code']}"} for r in left]
    return checks, summ, pack_only


_CMP_COLS = r"white|grey|gray|black|orange|silver|red|green|blue|clear|opal|frosted|brass|gold|yellow|pink|purple|cream|brown"


def ck_compare_table(signs, summ, reads=None, retailer="TK Maxx"):
    """One row per sign: code / colour / height, yours vs the pack's, with a match flag. reads = drawing-check results (optional)."""
    reads = reads or {}
    out = []
    for s, sm in zip(signs, summ):
        pk = sm.get("pk", [])
        rd = reads.get(s["title"], {})
        man_codes = sorted(set(TK_CODE_RE.findall((s["title"] + " " + s["text"]).upper()))) if retailer == "TK Maxx" else []
        pack_codes = sorted({r["code"] for r in pk})
        if retailer == "TK Maxx":
            if man_codes and pack_codes:
                cm = "Yes" if all(any(tk_code_same(p, m) for p in pack_codes) for m in man_codes) else "No"
            else:
                cm = "?"
        else:
            cm = "n/a"
        # colour
        mcol = sorted({w.lower().replace("gray", "grey") for w in re.findall(_CMP_COLS, s["text"], re.I)})
        pcol = sorted({w.lower().replace("gray", "grey") for r in pk for w in re.findall(_CMP_COLS, r["desc"], re.I)})
        if rd:
            mcol = sorted(set(mcol) | set(rd["man"]["colours"])) if not rd["man"]["colours"] else rd["man"]["colours"]
            pcol = sorted(set(pcol) | set(rd["pack"]["colours"]))
        colm = "?" if not (mcol and pcol) else ("Yes" if set(mcol) & set(pcol) else "No")
        # height: letter height / fixing height / stated sizes
        pH = sorted({int(n) for r in pk for n in re.findall(r"(\d{2,4})\s*mm", (r.get("size_name") or "") + " " + r["desc"])})
        mH = sorted(set(s.get("ffl", [])))
        if rd:
            mH = sorted(set(mH) | set(rd["man"]["letter"]) | set(rd["man"]["ffl"]))
            pH = sorted(set(pH) | set(rd["pack"]["letter"]) | set(rd["pack"]["ffl"]))
        hm = "?" if not (mH and pH) else ("Yes" if set(mH) & set(pH) else "No")
        out.append({"Sign": s["title"], "Code - yours": ", ".join(man_codes) or "-", "Code - pack": ", ".join(pack_codes) or "-", "Code match": cm,
                    "Colour - yours": ", ".join(mcol[:4]) or "-", "Colour - pack": ", ".join(pcol[:4]) or "-", "Colour match": colm,
                    "Height - yours": ", ".join(f"{h}" for h in mH[:4]) or "-", "Height - pack": ", ".join(f"{h}" for h in pH[:4]) or "-", "Height match": hm,
                    "Pages": ", ".join(f"p{p}" for p in s["pages"][:3])})
    return out


# ---- Revision Checker -------------------------------------------------------------------------
def _rv_norm(desc):
    return " ".join(sorted(ck_tokens(desc, True)))


def _rv_pos(p):
    # sheet numbers get re-issued between revisions (A432 -> A431), so positions are compared by view name only
    return sorted({re.sub(r"^A\d{3}\s*", "", re.sub(r"\s*x\d+\s*$", "", x).strip()).lower() for x in (p or "").split(";") if x.strip()})


def rv_diff_rows(old, new, retailer):
    """Pair rows old<->new, return list of change dicts."""
    import difflib
    old = [r for r in old if r.get("qty", 0) > 0 or r.get("desc")]
    new = [r for r in new if r.get("qty", 0) > 0 or r.get("desc")]
    pairs, uo, un = [], list(range(len(old))), list(range(len(new)))

    def take(fn, thr):
        cand = []
        for i in uo:
            for j in un:
                sc = fn(old[i], new[j])
                if sc >= thr:
                    cand.append((sc, i, j))
        for sc, i, j in sorted(cand, reverse=True):
            if i in uo and j in un:
                uo.remove(i); un.remove(j); pairs.append((i, j, sc))
    if retailer == "TK Maxx":
        take(lambda a, b: 1.0 if a["code"] and a["code"] == b["code"] and a["level"] == b["level"] else 0.0, 1.0)
    take(lambda a, b: 1.0 if _rv_norm(a["desc"]) == _rv_norm(b["desc"]) and a["level"] == b["level"] else 0.0, 1.0)
    take(lambda a, b: 0.99 if _rv_norm(a["desc"]) == _rv_norm(b["desc"]) else 0.0, 0.99)
    sim = lambda a, b: difflib.SequenceMatcher(None, a["desc"].lower(), b["desc"].lower()).ratio()
    take(sim, 0.85)
    take(sim, 0.6)
    out = []
    for i, j, sc in pairs:
        a, b = old[i], new[j]
        if a["qty"] != b["qty"]:
            out.append(("Quantity changed", b, f"{a['qty']}", f"{b['qty']}", a, "Update the 'No. required' on the sheet"))
        if _rv_pos(a["position"]) != _rv_pos(b["position"]) and (a["position"] or b["position"]):
            out.append(("Position changed", b, a["position"][:80] or "-", b["position"][:80] or "-", a, "Update the position / elevation on the sheet"))
        if (a["size_name"] or "") != (b["size_name"] or ""):
            out.append(("Size changed", b, a["size_name"] or "-", b["size_name"] or "-", a, "Update dimensions and artwork"))
        if re.sub(r"\s+", " ", a["desc"]).strip().lower() != re.sub(r"\s+", " ", b["desc"]).strip().lower():
            out.append(("Wording changed", b, a["desc"][:110], b["desc"][:110], a, "Check spec / artwork against the new wording"))
    for j in un:
        b = new[j]
        out.append(("Added", b, "-", f"{b['qty']} x {b['desc'][:100]}", None, "New manual sheet needed"))
    for i in uo:
        a = old[i]
        out.append(("Removed", a, f"{a['qty']} x {a['desc'][:100]}", "-", a, "Remove / mark the manual sheet"))
    return out


def _rv_sheets(pack_bytes):
    sc = ms_scan(pack_bytes)
    return {p["page"]: p for p in sc["pages"] if p["type"] not in ("error", "duplicate")}


def _rv_inkdiff(a, b, max_shift=20):
    """Ink that is in one sheet but not near the same place in the other. Lines the sheets up first (a re-issued drawing is often
    nudged a few pixels), then ignores 1px drift, so only genuine redraws show."""
    import numpy as np
    ia, ib = a < 200, b < 200
    h, w = ia.shape
    # best shift by FFT cross-correlation of the ink maps, limited to +/- max_shift px
    fa = np.fft.rfft2(ia.astype("float32")); fb = np.fft.rfft2(ib.astype("float32"))
    cc = np.fft.irfft2(fa * np.conj(fb), s=ia.shape)
    best, dy, dx = -1, 0, 0
    for y in list(range(0, max_shift + 1)) + list(range(h - max_shift, h)):
        for x in list(range(0, max_shift + 1)) + list(range(w - max_shift, w)):
            if cc[y, x] > best:
                best, dy, dx = cc[y, x], y, x
    dy = dy - h if dy > h // 2 else dy
    dx = dx - w if dx > w // 2 else dx
    ib = np.roll(np.roll(ib, dy, axis=0), dx, axis=1)
    def dil(m):
        out = m.copy()
        for oy in (-1, 0, 1):
            for ox in (-1, 0, 1):
                out |= np.roll(np.roll(m, oy, axis=0), ox, axis=1)
        return out
    d = (ib & ~dil(ia)) | (ia & ~dil(ib))
    m = max_shift                                          # edges exposed by the shift
    d[:m, :] = d[-m:, :] = False; d[:, :m] = d[:, -m:] = False
    return d


def _rv_pixdiff(old_bytes, new_bytes, pairs, scale=0.3, min_frac=0.0002):
    """pairs = [(old page, new page, label)]. Returns changed-page records with before/after PNG crops."""
    import numpy as np
    from PIL import ImageChops, ImageFilter
    o, n = pdfium.PdfDocument(old_bytes), pdfium.PdfDocument(new_bytes)
    out = []
    try:
        for po, pn, label in pairs:
            a_pg, b_pg = o[po - 1], n[pn - 1]
            try:
                A = a_pg.render(scale=scale).to_pil().convert("L")
                B = b_pg.render(scale=scale).to_pil().convert("L")
                sized = abs(A.size[0] - B.size[0]) > 4 or abs(A.size[1] - B.size[1]) > 4
                if sized:
                    B = B.resize(A.size)
                d = _rv_inkdiff(np.asarray(A), np.asarray(B))
                w = d.shape[1]
                d[:, int(w * 0.85):] = False          # title-block strip (rev table, dates, status stamps) - not sign content
                frac = float(d.mean())                # share of the sheet's pixels that changed (before padding for the crop)
                d = np.asarray(Image.fromarray((d * 255).astype("uint8")).filter(ImageFilter.MaxFilter(9))) > 0
                if frac < min_frac:
                    continue
                ys, xs = np.where(d)
                x0, x1, y0, y1 = xs.min(), xs.max(), ys.min(), ys.max()
                pad = 60
                bx = (max(0, x0 - pad) / scale, max(0, y0 - pad) / scale, min(w, x1 + pad) / scale, min(d.shape[0], y1 + pad) / scale)
                for_big = 0.9 if (bx[2] - bx[0]) < 900 else 0.5
                ia = crop_image(a_pg, bx, for_big).resize((int((bx[2] - bx[0]) * for_big), int((bx[3] - bx[1]) * for_big)))
                bb = crop_image(b_pg, bx, for_big)
                bufa, bufb = io.BytesIO(), io.BytesIO()
                ia.save(bufa, "PNG"); bb.save(bufb, "PNG")
                out.append({"label": label, "old": po, "new": pn, "frac": frac, "size_differs": sized,
                            "before": bufa.getvalue(), "after": bufb.getvalue()})
            finally:
                a_pg.close(); b_pg.close()
    finally:
        o.close(); n.close()
    return out


def _ck_revision(old_bytes, new_bytes, retailer, with_pixels=True, ver=""):
    po, pn = _ck_pack(old_bytes, retailer, False), _ck_pack(new_bytes, retailer, False)
    changes = rv_diff_rows(po["rows"], pn["rows"], retailer)
    sheets_old, sheets_new = {}, {}
    pages_old = {p["page"]: p for p in (po["pages"] or [])}
    pages_new = {p["page"]: p for p in (pn["pages"] or [])}
    def _keyed(pages):
        # M&S pages carry a drawing number + title; TK Maxx pages carry a sheet number ('A410'), sometimes shared by two pages
        out, seen = {}, Counter()
        for p in sorted(pages.values(), key=lambda p: p["page"]):
            if p.get("type") in ("error", "duplicate"):
                continue
            base = (p.get("dwg") or p.get("sheet") or "", "" if (p.get("dwg") or p.get("sheet")) else (p.get("title") or "").lower())
            if base == ("", ""):
                base = ("", f"page {p['page']}")
            seen[base] += 1
            out[base if seen[base] == 1 else (base[0], f"{base[1]}#{seen[base]}")] = p
        return out
    so, sn = _keyed(pages_old), _keyed(pages_new)
    # sheets renamed between issues ('Foodhall Internal Elevations' -> 'Proposed Ground Floor Foodhall Internal Elevations'):
    # pair the leftovers by kind and title similarity so they are compared, not reported as removed + added
    lo = [k for k in so if k not in sn]; ln = [k for k in sn if k not in so]
    for k in list(lo):
        best = max(ln, key=lambda j: (so[k].get("type") == sn[j].get("type"),
                                      difflib.SequenceMatcher(None, (so[k].get("title") or "").lower(), (sn[j].get("title") or "").lower()).ratio()),
                   default=None)
        if best is not None and so[k].get("type") == sn[best].get("type") and (
                len(lo) == 1 and len(ln) == 1 or
                difflib.SequenceMatcher(None, (so[k].get("title") or "").lower(), (sn[best].get("title") or "").lower()).ratio() >= 0.6):
            sn[k] = sn.pop(best); ln.remove(best); lo.remove(k)
    sheet_changes = []
    for k, p in sn.items():
        if k not in so:
            sheet_changes.append({"Change": "Sheet added", "Sheet": f"{k[0]} {p.get('title','') or p.get('type','')}".strip(), "Old page": "-", "New page": f"p{p['page']}"})
    for k, p in so.items():
        if k not in sn:
            sheet_changes.append({"Change": "Sheet removed", "Sheet": f"{k[0]} {p.get('title','') or p.get('type','')}".strip(), "Old page": f"p{p['page']}", "New page": "-"})
    pix = []
    if with_pixels and pages_old and pages_new:
        pairs = []
        for k, p in sn.items():
            if k in so and p.get("type") in ("plan", "elevation") and so[k].get("type") in ("plan", "elevation"):
                pairs.append((so[k]["page"], p["page"], f"{k[0]} {p.get('title','') or p.get('type','')} (p{p['page']})".strip()))
        pix = _rv_pixdiff(old_bytes, new_bytes, pairs)
        pix.sort(key=lambda r: -r["frac"])
    return {"changes": changes, "sheets": sheet_changes, "pix": pix, "old_rows": po["rows"], "new_rows": pn["rows"],
            "n_old": len(po["rows"]), "n_new": len(pn["rows"])}


ck_revision = st.cache_data(show_spinner="Comparing the two packs...")(_ck_revision)


# ---- A4 landscape report ----------------------------------------------------------------------
def ck_report_pdf(title, subtitle, metrics, sections, notes=(), status_cols=("Status",), images=()):
    """sections = [(heading, DataFrame, [col widths in mm or None])]; images = [(caption, before png, after png)]."""
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image as RLImage, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=11 * mm,
                            bottomMargin=13 * mm, title=title)
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], fontSize=17, leading=20, alignment=0, spaceAfter=2)
    sub = ParagraphStyle("sub", parent=ss["Normal"], fontSize=8.5, leading=11, textColor=colors.HexColor("#555555"))
    cell = ParagraphStyle("cell", parent=ss["Normal"], fontSize=7.4, leading=9)
    cellb = ParagraphStyle("cellb", parent=cell, fontName="Helvetica-Bold")
    head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    h2 = ParagraphStyle("h2", parent=ss["Normal"], fontSize=11, leading=14, fontName="Helvetica-Bold", spaceBefore=6, spaceAfter=4)
    P = lambda t, s=cell: Paragraph(escape(str(t if t is not None else "")), s)
    avail = landscape(A4)[0] - 24 * mm
    story = [Paragraph(escape(title), h1), Paragraph(escape(subtitle), sub), Spacer(1, 3 * mm)]
    if metrics:
        mt = Table([[Paragraph(f"<b>{escape(str(v))}</b><br/>{escape(k)}", ParagraphStyle("m", parent=cell, fontSize=9.5, leading=12, alignment=1))
                     for k, v in metrics]], hAlign="LEFT", colWidths=[34 * mm] * len(metrics))
        mt.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, colors.lightgrey), ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
                                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f4f5f7")),
                                ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        story += [mt, Spacer(1, 4 * mm)]
    pal = {"OK": "#d7f0dd", "Check": "#fff0c2", "Info": "#e6eefb", "Missing": "#f9d3d3", "Added": "#d7f0dd", "Removed": "#f9d3d3", "Yes": "#d7f0dd", "No": "#f9d3d3", "?": "#fff0c2"}
    for k, (heading, df, widths) in enumerate(sections):
        if k:
            story.append(Spacer(1, 4 * mm))
        story.append(Paragraph(escape(heading), h2))
        if df is None or df.empty:
            story.append(Paragraph("Nothing to report.", sub))
            continue
        cols = [c for c in df.columns if not c.startswith("_")]
        ws = [w * mm if w else None for w in (widths or [None] * len(cols))]
        fixed = sum(w for w in ws if w)
        free = [i for i, w in enumerate(ws) if not w]
        ws = [w if w else (avail - fixed) / max(len(free), 1) for w in ws]
        data = [[P(c, head) for c in cols]] + [[P(r[c], cellb if c in status_cols or c == cols[0] else cell) for c in cols] for _, r in df.iterrows()]
        t = Table(data, colWidths=ws, repeatRows=1)
        st_ = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2f3a4a")), ("GRID", (0, 0), (-1, -1), 0.4, colors.lightgrey),
               ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8f9fb")])]
        for c in cols:
            if c in status_cols or c == "Change":
                ci = cols.index(c)
                for n, (_, r) in enumerate(df.iterrows(), start=1):
                    if str(r[c]) in pal:
                        st_.append(("BACKGROUND", (ci, n), (ci, n), colors.HexColor(pal[str(r[c])])))
                    elif c == "Change":
                        st_.append(("BACKGROUND", (ci, n), (ci, n), colors.HexColor("#fff0c2")))
        t.setStyle(TableStyle(st_))
        story.append(t)
    if notes:
        story += [Spacer(1, 4 * mm), Paragraph("Notes", h2)]
        story += [Paragraph("&bull; " + escape(n), sub) for n in notes]
    if images:
        story += [PageBreak(), Paragraph("Drawings that changed (before | after)", h2)]
        from reportlab.lib.utils import ImageReader
        for cap, a, b in images:
            story.append(Paragraph(escape(cap), cellb))
            ims = []
            for png in (a, b):
                ir = ImageReader(io.BytesIO(png)); iw, ih = ir.getSize()
                w = avail / 2 - 3 * mm; h = min(ih * w / iw, 62 * mm); w = min(w, iw * h / ih)
                ims.append(RLImage(io.BytesIO(png), width=w, height=h))
            tb = Table([ims], colWidths=[avail / 2] * 2)
            tb.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.4, colors.lightgrey), ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.lightgrey),
                                    ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story += [tb, Spacer(1, 3 * mm)]

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(12 * mm, 7 * mm, title)
        canvas.drawRightString(landscape(A4)[0] - 12 * mm, 7 * mm, f"Page {d.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ---- Drawing check (optional, API): read height / FFL / colour / wording from BOTH drawings and compare -----------
DRAW_PROMPT = """You are reading one signage drawing for a sign-manufacturing check. The sign is: "{title}" (code {code}).
The image is {where}. Read ONLY what is drawn for this sign - dimensions, colours and wording on or next to it. On a client elevation the sign is the one that the boxed tag "{code}" points to with its leader line; ignore neighbouring signs and their dimensions.
Return ONLY a JSON object, no other text:
{{"letter_height_mm": [numbers], "overall_size_mm": [numbers], "fixing_height_mm": [numbers],
 "colours": ["white", ...], "wording": "text on the sign as you read it", "unclear": "anything you could not read or are unsure of"}}
letter_height_mm = height of the lettering (e.g. a vertical dimension beside the letters).
overall_size_mm = overall length / width / height of the sign itself.
fixing_height_mm = how high it is fixed from finished floor level (FFL / AFFL).
colours = the colour(s) the sign face / letters are shown in (use simple words: white, grey, black, orange, silver, red, green, blue, clear, opal, frosted, brass, gold, yellow).
Use empty lists / empty strings where the drawing does not show it. Never guess."""

_DRAW_COLS = r"white|grey|gray|black|orange|silver|red|green|blue|clear|opal|frosted|brass|gold|yellow|pink|purple|cream|brown"


def _ck_img_b64(png, maxw=1800):
    im = Image.open(io.BytesIO(png)).convert("RGB")
    if im.size[0] > maxw:
        im = im.resize((maxw, int(im.size[1] * maxw / im.size[0])))
    buf = io.BytesIO()
    im.save(buf, format="JPEG", quality=88)
    return base64.b64encode(buf.getvalue()).decode()


def ck_draw_read(client, model, png, title, code, where, reader=None):
    prompt = DRAW_PROMPT.format(title=title, code=code or "n/a", where=where)
    if reader is not None:
        text = reader(png, prompt)
    else:
        msg = client.messages.create(model=model, max_tokens=700, messages=[{"role": "user", "content": [
            {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": _ck_img_b64(png)}},
            {"type": "text", "text": prompt}]}])
        text = "".join(b.text for b in msg.content if getattr(b, "type", "") == "text")
    a, b = text.find("{"), text.rfind("}")
    d = json.loads(text[a:b + 1])
    num = lambda v: sorted({int(float(x)) for x in (v or []) if re.fullmatch(r"\d+(?:\.\d+)?", str(x).strip())})
    return {"letter": num(d.get("letter_height_mm")), "size": num(d.get("overall_size_mm")), "ffl": num(d.get("fixing_height_mm")),
            "colours": sorted({c.lower().replace("gray", "grey") for c in (d.get("colours") or []) if re.fullmatch(_DRAW_COLS, str(c).lower())}),
            "wording": str(d.get("wording") or ""), "unclear": str(d.get("unclear") or "")}


def ck_draw_merge(reads):
    out = {"letter": [], "size": [], "ffl": [], "colours": [], "wording": "", "unclear": ""}
    for r in reads:
        for k in ("letter", "size", "ffl", "colours"):
            out[k] = sorted(set(out[k]) | set(r[k]))
        out["wording"] = out["wording"] or r["wording"]
        out["unclear"] = (out["unclear"] + " " + r["unclear"]).strip()
    return out


def ck_draw_compare(man, pack):
    """Rows: Item, Manual drawing, Pack drawing, Status, Note."""
    rows = []
    for key, label, unit in (("letter", "Letter height", "mm"), ("ffl", "Fixing height (FFL)", "mm"), ("size", "Overall size", "mm")):
        a, b = man[key], pack[key]
        if not a and not b:
            continue
        if not a or not b:
            rows.append((label, ", ".join(f"{x}{unit}" for x in a) or "not read", ", ".join(f"{x}{unit}" for x in b) or "not read",
                         "Not read", "Only one drawing shows it - look at both"))
        else:
            ok = bool(set(a) & set(b))
            rows.append((label, ", ".join(f"{x}{unit}" for x in a), ", ".join(f"{x}{unit}" for x in b), "OK" if ok else "Check",
                         "" if ok else "Heights/sizes differ between your sheet and the pack"))
    a, b = man["colours"], pack["colours"]
    if a or b:
        if not a or not b:
            rows.append(("Colour", ", ".join(a) or "not read", ", ".join(b) or "not read", "Not read", "Only one drawing shows a colour"))
        else:
            ok = bool(set(a) & set(b))
            rows.append(("Colour", ", ".join(a), ", ".join(b), "OK" if ok else "Check", "" if ok else "Colours differ"))
    ta, tb = ck_tokens(man["wording"], True), ck_tokens(pack["wording"], True)
    if ta and tb:
        ok = len(ta & tb) / min(len(ta), len(tb)) >= 0.6
        rows.append(("Wording", man["wording"][:70], pack["wording"][:70], "OK" if ok else "Check", "" if ok else "Wording differs"))
    return rows


def _ck_render_page_png(pdf_bytes, pno, scale=1.6):
    doc = pdfium.PdfDocument(pdf_bytes)
    try:
        pg = doc[pno - 1]
        try:
            im = pg.render(scale=scale).to_pil().convert("RGB")
        finally:
            pg.close()
    finally:
        doc.close()
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def ck_draw_check(client, model, signs, checks, pack, manual_bytes, only=None, reader=None, progress=None):
    """Run the drawing read for the chosen signs. Returns [{sign, item, manual, pack, status, note}]."""
    out = []
    reads_out = {}
    by_sign = {}
    for r in checks:
        by_sign.setdefault(r["Sign"], r)
    chosen = [s for s in signs if (only is None or s["title"] in only)]
    for n, s in enumerate(chosen):
        r = by_sign.get(s["title"])
        keys = r["_keys"] if r else []
        snips = [sn for k in keys for sn in pack["snips"].get(k, [])]
        snips = sorted(snips, key=lambda x: 0 if x.get("kind") == "elevation" else 1)[:2]
        if progress:
            progress(n / max(len(chosen), 1), s["title"])
        if not snips:
            out.append({"sign": s["title"], "item": "-", "manual": "", "pack": "no pack drawing found for this sign", "status": "Not read", "note": ""})
            continue
        try:
            code = ", ".join(sorted({k.split("|")[-1] for k in keys}))[:40]
            man = ck_draw_read(client, model, _ck_render_page_png(manual_bytes, s["pages"][0]), s["title"], code,
                               "the manufacturer's sign sheet (detail drawing plus an elevation showing where it fits)", reader)
            pk = ck_draw_merge([ck_draw_read(client, model, sn["png"], s["title"], code, "a crop of the client's elevation drawing around this sign", reader)
                                for sn in snips])
        except Exception as e:
            out.append({"sign": s["title"], "item": "-", "manual": "", "pack": f"could not read: {e}"[:90], "status": "Not read", "note": ""})
            continue
        reads_out[s["title"]] = {"man": man, "pack": pk}
        rows = ck_draw_compare(man, pk)
        if not rows:
            out.append({"sign": s["title"], "item": "-", "manual": "", "pack": "nothing readable on either drawing", "status": "Not read", "note": ""})
        for it, a, b, st_, note in rows:
            out.append({"sign": s["title"], "item": it, "manual": a, "pack": b, "status": st_, "note": note})
    return out, reads_out


# ---- Screens: Manual Checker + Revision Checker ----------------------------------------------
def _ck_manual_png(manual_bytes, pno, scale=0.9):
    doc = pdfium.PdfDocument(manual_bytes)
    try:
        pg = doc[pno - 1]
        try:
            im = pg.render(scale=scale).to_pil().convert("RGB")
        finally:
            pg.close()
    finally:
        doc.close()
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def _ck_df(rows, cols):
    return pd.DataFrame([{c: r.get(c, "") for c in cols} for r in rows], columns=cols)


def manual_checker(retailer):
    st.subheader(f"Manual Checker - {retailer}")
    st.caption("Checks your finished sign manual against the pack the client supplied, sign by sign: is it in the pack, "
               "does the quantity agree, do the sizes agree, and does the pack itself flag it as unsettled.")
    u1, u2 = st.columns(2)
    mf = u1.file_uploader("Your manual (PDF, Geetee sheets)", type=["pdf"], key=f"mc_manual_{retailer}")
    pf = u2.file_uploader(f"{retailer} pack from the client (PDF)", type=["pdf"], key=f"mc_pack_{retailer}")
    if not (mf and pf):
        st.info("Upload both files. Internal signage only - sheets such as plans and elevations in your manual are read as layout, not as signs.")
        return
    mb, pb = mf.getvalue(), pf.getvalue()
    try:
        pages = read_manual(mb, CK_VER)
        signs, layout = manual_signs(pages)
        pack = ck_pack(pb, retailer, True, CK_VER)
    except Exception as e:
        st.error(f"Couldn't read one of those files: {e}")
        return
    if not signs:
        st.error("No Geetee sign sheets found in the manual (looking for the DESCRIPTION / DRAWING NUMBER title block).")
        return
    if not pack["rows"]:
        st.error("Nothing found in the pack. Check it's the right drawing pack for this retailer.")
        return
    for w in pack["warnings"]:
        st.warning(w)
    checks, summ, pack_only = ck_check(signs, pack["rows"], retailer)
    df = _ck_df(checks, ["Sign", "Dwg", "Check", "Manual says", "Pack says", "Status", "Note", "Manual page", "Pack page"])
    po = _ck_df(pack_only, ["Sign", "Code", "Qty", "Where", "Pack page", "Status", "Note"])
    n_ok = sum(1 for s in summ if s["status"] == "OK")
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Signs in manual", len(summ))
    c2.metric("All OK", n_ok)
    c3.metric("To check", len(summ) - n_ok)
    c4.metric("In pack, not in manual", len(po))
    st.caption(f"Pack read in {'signage-key' if pack['mode'] == 'key' else 'callout'} mode: {len(pack['rows'])} rows. "
               f"{len(layout)} layout sheet(s) in your manual were skipped.")

    dkey = f"mc_draw_{retailer}_{hashlib.md5(mb + pb).hexdigest()}"
    reads = st.session_state.get(dkey + "_reads", {})
    cmp_rows = ck_compare_table(signs, summ, reads, retailer)
    cmp_df = pd.DataFrame(cmp_rows)
    t0, t1, t2, t3 = st.tabs(["Side by side", "Checks", "In the pack, not in your manual", "Manual sheets read"])
    with t0:
        st.caption("Yours vs the pack for each sign. 'Yes' = they agree, 'No' = they differ, '?' = one side couldn't be read from the text - "
                   "run the Drawing check below to read it from the drawings. Heights are mm (fixing height and letter height).")
        only_bad = st.checkbox("Only show signs with a 'No' or '?'", value=False, key=f"mc_cmp_bad_{retailer}")
        vw = cmp_df
        if only_bad:
            vw = cmp_df[(cmp_df[["Code match", "Colour match", "Height match"]] != "Yes").any(axis=1)]
        st.dataframe(vw, use_container_width=True, hide_index=True)
    with t1:
        show = st.multiselect("Show", ["OK", "Check", "Info"], default=["OK", "Check", "Info"], key=f"mc_filter_{retailer}")
        only_issue = st.checkbox("Hide the 'In the pack' OK lines for signs that are otherwise fine", value=False, key=f"mc_hide_{retailer}")
        view = df[df["Status"].isin(show)].reset_index(drop=True)
        if only_issue:
            bad = set(df.loc[df["Status"] == "Check", "Sign"])
            view = view[view["Sign"].isin(bad)].reset_index(drop=True)
        st.dataframe(view, use_container_width=True, hide_index=True)
        if not view.empty:
            allv = [r for r in checks if r["Status"] in show]
            i = st.selectbox("Look at one", range(len(allv)), key=f"mc_pick_{retailer}",
                             format_func=lambda k: f"{allv[k]['Sign']} - {allv[k]['Check']} ({allv[k]['Status']})")
            r = allv[i]
            a, b = st.columns(2)
            a.image(_ck_manual_png(mb, r["_mp"]), caption=f"Your manual, p{r['_mp']}")
            snip = next((s for k in r["_keys"] for s in pack["snips"].get(k, [])), None)
            if snip is not None:
                b.image(snip["png"] if isinstance(snip, dict) else snip, caption="Pack drawing / callout")
            else:
                b.info(f"Pack: {r['Pack says'] or 'nothing matched'}" + (f"  ({r['Pack page']})" if r["Pack page"] else ""))
            if r["Note"]:
                st.info(r["Note"])
    with t2:
        st.caption("Pack items that no sheet in your manual seems to cover - new signs, existing/retained items, client-supplied graphics, or noise.")
        st.dataframe(po, use_container_width=True, hide_index=True)
    with t3:
        sm = pd.DataFrame([{"Pages": s["pages"], "Dwg": s["dwg"], "Sign": s["sign"], "Pack rows": s["n_pack"],
                            "Status": s["status"], "Last revision": s["rev"]} for s in summ])
        st.dataframe(sm, use_container_width=True, hide_index=True)
        if layout:
            st.caption("Skipped as layout: " + "; ".join(f"p{p['page']} {p['title']}" for p in layout))

    with st.expander("Drawing check - read heights, FFL, colour and wording from both drawings (uses the API)"):
        st.caption("Reads your sign sheet and the pack's elevation around each sign, then compares letter height, fixing height, overall size, colour "
                   "and wording. It is a prompt to look, not a verdict - open the sign in the viewer above to confirm by eye.")
        dkey_api = get_api_key() or st.text_input("Anthropic API key", type="password", key=f"mc_draw_key_{retailer}")
        dmodel = st.selectbox("Model", ["claude-sonnet-5-5", "claude-haiku-5-5"], key=f"mc_draw_model_{retailer}")
        names = [s["title"] for s in signs]
        pick = st.multiselect("Signs to read", names, default=names, key=f"mc_draw_pick_{retailer}")
        if st.button(f"Run drawing check on {len(pick)} sign(s)", key=f"mc_draw_go_{retailer}", disabled=not (dkey_api and pick)):
            import anthropic
            bar = st.progress(0.0)
            res_d, reads_d = ck_draw_check(anthropic.Anthropic(api_key=dkey_api), dmodel, signs, checks, pack, mb, set(pick),
                                  progress=lambda f, t: bar.progress(min(f, 1.0), text=t))
            bar.empty()
            st.session_state[dkey] = res_d
            st.session_state[dkey + "_reads"] = reads_d
            st.rerun()
        if not dkey_api:
            st.info("Set ANTHROPIC_API_KEY (or paste a key here) to run this.")
    draw = st.session_state.get(dkey)
    if draw:
        ddf = pd.DataFrame(draw).rename(columns={"sign": "Sign", "item": "Item", "manual": "Your sheet", "pack": "Pack drawing", "status": "Status", "note": "Note"})
        st.subheader("Drawing check")
        st.dataframe(ddf, use_container_width=True, hide_index=True)

    sec = [("Yours vs the pack - code, colour, height", cmp_df, [38, 22, 22, 12, 22, 22, 15, 16, 16, 15, None]),
           ("Checks needing attention", df[df["Status"] == "Check"].drop(columns=["Dwg"]), [40, 24, 52, 62, 16, None, 14, 14]),
           ("Pack items not covered by a manual sheet", po, [70, 20, 10, 50, 16, 16, None])]
    if draw:
        sec.append(("Drawing check (read from the drawings by the vision model - confirm by eye)", ddf[ddf["Status"] != "OK"], [40, 30, 50, 50, 16, None]))
    pdf = ck_report_pdf(f"Manual check - {retailer}", f"{mf.name}  checked against  {pf.name}  (internal signage only)",
                        [("Signs in manual", len(summ)), ("All OK", n_ok), ("To check", len(summ) - n_ok), ("Pack, not in manual", len(po))],
                        [(h, d, w) for (h, d, w) in sec],
                        notes=["Signs are matched to the pack by name (and by code for TK Maxx), so a Check line may be a naming difference rather than a real error.",
                               "Quantities are compared to what your sheet states ('N No. required'); sheets that state none are not compared.",
                               "Layout-only sheets (plans, elevations) in the manual are not treated as signs.",
                               "Side-by-side table: Yes = agree, No = differ, ? = one side could not be read from the text."],
                        status_cols=("Status", "Code match", "Colour match", "Height match"))
    d1, d2 = st.columns(2)
    d1.download_button("Download check sheet (PDF, A4 landscape)", pdf, "manual_check.pdf", "application/pdf", key=f"mc_dl_{retailer}")
    d2.download_button("Download checks (CSV)", df.to_csv(index=False).encode("utf-8-sig"), "manual_check.csv", "text/csv", key=f"mc_csv_{retailer}")


def _rv_impact(change, signs, idf):
    if not signs:
        return change[5]
    r = change[1]
    sc = max(((ck_score(s, r, idf), s) for s in signs), key=lambda x: x[0], default=(0, None))
    if sc[0] >= 0.5:
        s = sc[1]
        return f"{change[5].split(' on ')[0]} - manual p{', p'.join(str(p) for p in s['pages'])} ({s['dwg']})"
    return "No manual sheet found - " + ("new sheet needed" if change[0] == "Added" else "check whether one is needed")


def revision_checker(retailer):
    st.subheader(f"Revision Checker - {retailer}")
    st.caption("Compares the pack you've worked from with the revised pack: signs added, removed or changed, drawings that have been "
               "redrawn, and (if you add your manual) which manual sheets are affected.")
    u1, u2 = st.columns(2)
    of = u1.file_uploader("Pack you've worked from (PDF)", type=["pdf"], key=f"rc_old_{retailer}")
    nf = u2.file_uploader("Revised pack (PDF)", type=["pdf"], key=f"rc_new_{retailer}")
    mf = st.file_uploader("Optional: your manual (PDF) - adds 'which sheet to update'", type=["pdf"], key=f"rc_man_{retailer}")
    vis = st.checkbox("Also compare the drawings visually (finds changes the text doesn't show)", value=True, key=f"rc_vis_{retailer}")
    if not (of and nf):
        st.info("Upload the old and the revised pack.")
        return
    ob, nb = of.getvalue(), nf.getvalue()
    if hashlib.md5(ob).hexdigest() == hashlib.md5(nb).hexdigest():
        st.success("These two files are identical - nothing has changed.")
        return
    try:
        rv = ck_revision(ob, nb, retailer, vis, CK_VER)
    except Exception as e:
        st.error(f"Couldn't compare those files: {e}")
        return
    signs = []
    if mf:
        try:
            signs, _ = manual_signs(read_manual(mf.getvalue(), CK_VER))
        except Exception as e:
            st.warning(f"Couldn't read the manual, so manual impact isn't shown: {e}")
    idf = ck_idf(signs, rv["old_rows"] + rv["new_rows"]) if signs else None
    chg = rv["changes"]
    rows = [{"Change": c[0], "Sign": c[1]["desc"][:90], "Code": c[1]["code"], "Was": c[2], "Now": c[3],
             "Where": (c[1]["position"] or c[1]["level"])[:60],
             "Manual impact": _rv_impact(c, signs, idf) if signs else c[5]} for c in chg]
    df = pd.DataFrame(rows, columns=["Change", "Sign", "Code", "Was", "Now", "Where", "Manual impact"])
    sh = pd.DataFrame(rv["sheets"], columns=["Change", "Sheet", "Old page", "New page"])
    pix = rv["pix"]
    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("Sign changes", len(df))
    c2.metric("Added", int((df["Change"] == "Added").sum()))
    c3.metric("Removed", int((df["Change"] == "Removed").sum()))
    c4.metric("Amended", int((~df["Change"].isin(["Added", "Removed"])).sum()))
    c5.metric("Drawings redrawn", len(pix))
    st.caption(f"Old pack: {rv['n_old']} rows. Revised pack: {rv['n_new']} rows. "
               "Sign changes compare the schedules read from each pack; 'Drawings redrawn' compares the sheets pixel by pixel "
               "(the title-block strip is ignored).")
    t1, t2, t3 = st.tabs(["Sign changes", "Sheets added / removed", "Drawings that changed"])
    with t1:
        kinds = sorted(df["Change"].unique())
        pick = st.multiselect("Change types", kinds, default=kinds, key=f"rc_filter_{retailer}")
        st.dataframe(df[df["Change"].isin(pick)].reset_index(drop=True), use_container_width=True, hide_index=True)
    with t2:
        st.dataframe(sh, use_container_width=True, hide_index=True)
    with t3:
        if not vis:
            st.info("Visual comparison is switched off.")
        elif not pix:
            st.success("No sheet shows a visible change outside the title block.")
        for r in pix[:25]:
            st.markdown(f"**{r['label']}**  -  old p{r['old']} / new p{r['new']}  ({r['frac'] * 100:.2f}% of the sheet changed"
                        + (", sheet size differs - treat with care" if r["size_differs"] else "") + ")")
            a, b = st.columns(2)
            a.image(r["before"], caption="Before")
            b.image(r["after"], caption="After")
    pdf = ck_report_pdf(f"Revision check - {retailer}", f"{of.name}  ->  {nf.name}",
                        [("Sign changes", len(df)), ("Added", int((df["Change"] == "Added").sum())), ("Removed", int((df["Change"] == "Removed").sum())),
                         ("Amended", int((~df["Change"].isin(["Added", "Removed"])).sum())), ("Drawings redrawn", len(pix))],
                        [("Sign changes", df, [24, 62, 14, 46, 46, 30, None]), ("Sheets added / removed", sh, [30, 100, 25, 25])],
                        notes=["Changes come from comparing the schedule read from each pack, so renamed signs can appear as Removed + Added.",
                               "Drawings redrawn are found by comparing rendered sheets; the right-hand title-block strip is ignored."],
                        status_cols=("Change",), images=[(f"{r['label']}  (old p{r['old']} / new p{r['new']})", r["before"], r["after"]) for r in pix[:12]])
    d1, d2 = st.columns(2)
    d1.download_button("Download change list (PDF, A4 landscape)", pdf, "revision_changes.pdf", "application/pdf", key=f"rc_dl_{retailer}")
    d2.download_button("Download changes (CSV)", df.to_csv(index=False).encode("utf-8-sig"), "revision_changes.csv", "text/csv", key=f"rc_csv_{retailer}")


st.title("Store Signage Tools")
tool = st.radio("Tool", ["Signage Extractor", "Manual Checker", "Revision Checker"], horizontal=True,
                key="tool", label_visibility="collapsed")
retailer = st.radio("Retailer", ["TK Maxx", "M&S"], horizontal=True, key="retailer", label_visibility="collapsed")
if tool == "Signage Extractor":
    if retailer == "TK Maxx":
        tk_mode()
    else:
        ms_mode()
elif tool == "Manual Checker":
    manual_checker(retailer)
else:
    revision_checker(retailer)
